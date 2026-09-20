from pathlib import Path

from lxml import etree
from pptx import Presentation
from scripts.normalize_ooxml import normalize
from scripts.validate_ooxml import validate


def test_python_pptx_package_passes(tmp_path: Path) -> None:
    path = tmp_path / "valid.pptx"
    Presentation().save(path)
    assert validate(path)["passed"]


def test_normalizer_repairs_presentation_order(tmp_path: Path) -> None:
    valid = tmp_path / "valid.pptx"
    broken = tmp_path / "broken.pptx"
    repaired = tmp_path / "repaired.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(valid)

    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(valid) as source, ZipFile(broken, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "ppt/presentation.xml":
                root = etree.fromstring(data)
                children = list(root)
                slide_ids = next(
                    child for child in children if etree.QName(child).localname == "sldIdLst"
                )
                slide_size = next(
                    child for child in children if etree.QName(child).localname == "sldSz"
                )
                root.remove(slide_ids)
                root.insert(root.index(slide_size) + 1, slide_ids)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, data)

    assert not validate(broken)["passed"]
    normalize(broken, repaired)
    assert validate(repaired)["passed"]
