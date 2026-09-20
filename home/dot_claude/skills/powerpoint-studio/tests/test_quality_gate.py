from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from scripts.quality_gate import audit


def save_clean_deck(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(6), Inches(0.7))
    run = title.text_frame.paragraphs[0].add_run()
    run.text = "A clear conclusion"
    run.font.size = Pt(30)
    slide.shapes.add_shape(
        1,
        Inches(0.8),
        Inches(1.7),
        Inches(5.0),
        Inches(3.0),
    )
    presentation.save(path)


def test_clean_deck_passes(tmp_path: Path) -> None:
    deck = tmp_path / "clean.pptx"
    save_clean_deck(deck)
    report = audit(deck)
    assert report["summary"]["errors"] == 0


def test_detects_out_of_bounds_and_placeholder(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        presentation.slide_width - Inches(0.2),
        Inches(1),
        Inches(2),
        Inches(1),
    )
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "{{TITLE}}"
    run.font.size = Pt(24)
    deck = tmp_path / "broken.pptx"
    presentation.save(deck)
    report = audit(deck)
    categories = {item["category"] for item in report["findings"]}
    assert {"bounds", "placeholder"} <= categories


def test_overlapping_boxes_without_overlapping_text_do_not_fail(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    top = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1.2))
    top.text_frame.paragraphs[0].text = "Title"
    top.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    bottom = slide.shapes.add_textbox(Inches(1), Inches(1.9), Inches(6), Inches(1.2))
    bottom.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    bottom.text_frame.paragraphs[0].text = "Subtitle"
    bottom.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    output = tmp_path / "boxes.pptx"
    presentation.save(output)
    report = audit(output, {"text_only_allowed_slides": [1]})
    assert not any(finding["category"] == "overlap" for finding in report["findings"])


def test_rule_crossing_text_fails(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(0.8))
    text.text_frame.paragraphs[0].text = "Rule collision"
    text.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    slide.shapes.add_shape(1, Inches(1), Inches(1.18), Inches(5), Inches(0.03))
    output = tmp_path / "crossing-rule.pptx"
    presentation.save(output)
    report = audit(output, {"text_only_allowed_slides": [1]})
    assert any(
        finding["category"] == "overlap" and "Rule" in finding["message"]
        for finding in report["findings"]
    )


def test_conservative_wrapping_detects_text_overflow(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2.5), Inches(0.3))
    text.text_frame.margin_top = 0
    text.text_frame.margin_bottom = 0
    text.text_frame.paragraphs[0].text = "Controlled variable"
    run = text.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(16)
    run.font.bold = True
    run._r.get_or_add_rPr().set("spc", "80")
    output = tmp_path / "wrapped.pptx"
    presentation.save(output)
    report = audit(output, {"text_only_allowed_slides": [1]})
    assert any(finding["category"] == "text-fit" for finding in report["findings"])


def test_label_text_escaping_pill_fails(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    pill = slide.shapes.add_shape(5, Inches(1), Inches(1), Inches(1.4), Inches(0.34))
    pill.name = "label-pill"
    label = slide.shapes.add_textbox(Inches(1.05), Inches(1.15), Inches(1.3), Inches(0.3))
    label.name = "label-pill-text"
    label.text_frame.margin_top = 0
    label.text_frame.margin_bottom = 0
    label.text_frame.paragraphs[0].text = "INTRODUCTION"
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    output = tmp_path / "escaping-pill.pptx"
    presentation.save(output)
    report = audit(output, {"text_only_allowed_slides": [1]})
    assert any(
        finding["category"] == "containment" and finding["severity"] == "error"
        for finding in report["findings"]
    )


def test_tiny_evidence_screenshot_fails(tmp_path: Path) -> None:
    image = tmp_path / "evidence.png"
    Image.new("RGB", (800, 600), "white").save(image)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(1.2), Inches(0.5))
    picture.name = "evidence-screenshot"
    output = tmp_path / "tiny-evidence.pptx"
    presentation.save(output)
    report = audit(output, {"text_only_allowed_slides": [1]})
    assert any(
        finding["category"] == "asset" and "too small" in finding["message"]
        for finding in report["findings"]
    )
