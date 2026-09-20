#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
import math
import re
import subprocess
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

EMU_PER_INCH = 914_400
PLACEHOLDER_PATTERNS = (
    re.compile(r"\{\{[^}]+\}\}"),
    re.compile(r"\b(?:todo|tbd|lorem ipsum|sample text|placeholder)\b", re.I),
    re.compile(r"\[(?:必填|要修正|仮|TODO)\]", re.I),
)
VISUAL_SHAPE_TYPES = {
    MSO_SHAPE_TYPE.PICTURE,
    MSO_SHAPE_TYPE.CHART,
    MSO_SHAPE_TYPE.TABLE,
    MSO_SHAPE_TYPE.MEDIA,
}
MICROCOPY_NAME = re.compile(
    r"(section|footer|page|label|caption|legend|note|badge|owner|number|eyebrow|source|unit)",
    re.I,
)
CONTAINMENT_NAME = re.compile(r"(pill|badge|tag|chip|label|callout)", re.I)
EVIDENCE_IMAGE_NAME = re.compile(r"(evidence|screenshot|notebook|source-image)", re.I)


@dataclass
class Finding:
    severity: str
    slide: int
    category: str
    message: str
    shape: str | None = None


def inches(value: int) -> float:
    return value / EMU_PER_INCH


def box(shape: Any) -> tuple[float, float, float, float]:
    return (
        inches(shape.left),
        inches(shape.top),
        inches(shape.width),
        inches(shape.height),
    )


def area(bounds: tuple[float, float, float, float]) -> float:
    return max(0.0, bounds[2]) * max(0.0, bounds[3])


def intersection_area(
    a: tuple[float, float, float, float],
    b: tuple[float, float, float, float],
) -> float:
    left = max(a[0], b[0])
    top = max(a[1], b[1])
    right = min(a[0] + a[2], b[0] + b[2])
    bottom = min(a[1] + a[3], b[1] + b[3])
    return max(0.0, right - left) * max(0.0, bottom - top)


def text_of(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def paragraph_font_size(paragraph: Any, default: float = 18.0) -> float:
    sizes = [
        run.font.size.pt
        for run in paragraph.runs
        if run.font.size is not None and run.font.size.pt > 0
    ]
    if sizes:
        return max(sizes)
    if paragraph.font.size is not None:
        return paragraph.font.size.pt
    return default


def paragraph_font_name(paragraph: Any, default: str = "DejaVu Sans") -> str:
    for run in paragraph.runs:
        if run.font.name:
            return run.font.name
    return default


def paragraph_is_bold(paragraph: Any) -> bool:
    return any(run.font.bold is True for run in paragraph.runs)


def paragraph_char_spacing_pt(paragraph: Any) -> float:
    values: list[float] = []
    for run in paragraph.runs:
        properties = run._r.rPr
        if properties is None:
            continue
        raw = properties.get("spc")
        if raw:
            try:
                values.append(int(raw) / 100)
            except ValueError:
                continue
    return max(values, default=0.0)


def display_width_units(text: str) -> float:
    units = 0.0
    for char in text:
        code = ord(char)
        if char.isspace():
            units += 0.32
        elif code >= 0x2E80:
            units += 1.0
        elif char.isupper():
            units += 0.72
        elif char in "ilI.,:;!'|":
            units += 0.28
        else:
            units += 0.55
    return units


@lru_cache(maxsize=128)
def font_path(font_name: str, bold: bool) -> str | None:
    query = f"{font_name}:style={'Bold' if bold else 'Regular'}"
    try:
        result = subprocess.run(
            ["fc-match", "-f", "%{file}", query],
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return None
    candidate = result.stdout.strip()
    return candidate or None


@lru_cache(maxsize=512)
def measured_text_width_pt(
    text: str,
    font_name: str,
    font_size_pt: float,
    bold: bool,
    char_spacing_pt: float,
) -> float:
    path = font_path(font_name, bold)
    if path:
        try:
            font = ImageFont.truetype(path, max(1, round(font_size_pt * 96 / 72)))
            base_width = font.getlength(text) * 72 / 96
            return base_width + max(0, len(text) - 1) * char_spacing_pt
        except OSError:
            pass
    return display_width_units(text) * font_size_pt + max(0, len(text) - 1) * char_spacing_pt


def wrap_tokens(text: str) -> list[str]:
    if re.search(r"[\u2e80-\uffff]", text):
        return list(text)
    return re.findall(r"\S+\s*", text) or [text]


def contains_cjk(text: str) -> bool:
    return bool(re.search(r"[\u2e80-\uffff]", text))


def wrapped_line_count(
    text: str,
    width_pt: float,
    font_name: str,
    font_size_pt: float,
    bold: bool,
    char_spacing_pt: float,
) -> int:
    lines = 0
    for explicit_line in text.splitlines() or [""]:
        tokens = wrap_tokens(explicit_line)
        width_limit = width_pt * (1.08 if contains_cjk(explicit_line) else 1.01)
        current = ""
        line_count = 1
        for token in tokens:
            candidate = current + token
            if (
                not current
                or measured_text_width_pt(
                    candidate.rstrip(),
                    font_name,
                    font_size_pt,
                    bold,
                    char_spacing_pt,
                )
                <= width_limit
            ):
                current = candidate
                continue
            line_count += 1
            current = token.lstrip()
            if (
                measured_text_width_pt(
                    current,
                    font_name,
                    font_size_pt,
                    bold,
                    char_spacing_pt,
                )
                > width_limit
            ):
                line_count += max(
                    0,
                    math.ceil(
                        measured_text_width_pt(
                            current,
                            font_name,
                            font_size_pt,
                            bold,
                            char_spacing_pt,
                        )
                        / width_limit
                    )
                    - 1,
                )
        lines += line_count
    return max(1, lines)


def wrapped_line_widths_pt(
    text: str,
    width_pt: float,
    font_name: str,
    font_size_pt: float,
    bold: bool,
    char_spacing_pt: float,
) -> list[float]:
    widths: list[float] = []
    for explicit_line in text.splitlines() or [""]:
        tokens = wrap_tokens(explicit_line)
        width_limit = width_pt * (1.08 if contains_cjk(explicit_line) else 1.01)
        current = ""
        for token in tokens:
            candidate = current + token
            candidate_width = measured_text_width_pt(
                candidate.rstrip(),
                font_name,
                font_size_pt,
                bold,
                char_spacing_pt,
            )
            if not current or candidate_width <= width_limit:
                current = candidate
                continue
            widths.append(
                min(
                    width_pt,
                    measured_text_width_pt(
                        current.rstrip(),
                        font_name,
                        font_size_pt,
                        bold,
                        char_spacing_pt,
                    ),
                )
            )
            current = token.lstrip()
        widths.append(
            min(
                width_pt,
                measured_text_width_pt(
                    current.rstrip(),
                    font_name,
                    font_size_pt,
                    bold,
                    char_spacing_pt,
                ),
            )
        )
    return widths or [0.0]


def estimated_text_height_pt(shape: Any) -> tuple[float, float]:
    frame = shape.text_frame
    width_pt = max(
        1.0,
        inches(shape.width) * 72 - inches(frame.margin_left + frame.margin_right) * 72,
    )
    available_height_pt = max(
        1.0,
        inches(shape.height) * 72 - inches(frame.margin_top + frame.margin_bottom) * 72,
    )
    required = 0.0
    for paragraph in frame.paragraphs:
        size = paragraph_font_size(paragraph)
        font_name = paragraph_font_name(paragraph)
        bold = paragraph_is_bold(paragraph)
        char_spacing = paragraph_char_spacing_pt(paragraph)
        text = paragraph.text or " "
        lines = wrapped_line_count(
            text,
            width_pt,
            font_name,
            size,
            bold,
            char_spacing,
        )
        line_height = size * 1.22
        required += lines * line_height
        if paragraph.space_before:
            required += paragraph.space_before.pt
        if paragraph.space_after:
            required += paragraph.space_after.pt
    return required, available_height_pt


def text_ink_box(shape: Any) -> tuple[float, float, float, float]:
    x, y, width, height = box(shape)
    frame = shape.text_frame
    left = inches(frame.margin_left)
    right = inches(frame.margin_right)
    top = inches(frame.margin_top)
    bottom = inches(frame.margin_bottom)
    content_width = max(0.001, width - left - right)
    content_height = max(0.001, height - top - bottom)
    required_pt, _ = estimated_text_height_pt(shape)
    ink_height = min(content_height, required_pt / 72)
    paragraph_widths: list[float] = []
    alignments: set[Any] = set()
    for paragraph in frame.paragraphs:
        size = paragraph_font_size(paragraph)
        paragraph_widths.extend(
            wrapped_line_widths_pt(
                paragraph.text or " ",
                content_width * 72,
                paragraph_font_name(paragraph),
                size,
                paragraph_is_bold(paragraph),
                paragraph_char_spacing_pt(paragraph),
            )
        )
        alignments.add(paragraph.alignment)
    ink_width = min(content_width, max(paragraph_widths, default=content_width * 72) / 72)
    alignment = next(iter(alignments)) if len(alignments) == 1 else None
    if alignment == PP_ALIGN.CENTER:
        ink_x = x + left + (content_width - ink_width) / 2
    elif alignment == PP_ALIGN.RIGHT:
        ink_x = x + width - right - ink_width
    else:
        ink_x = x + left
    anchor = frame.vertical_anchor
    if anchor == MSO_ANCHOR.BOTTOM:
        ink_y = y + height - bottom - ink_height
    elif anchor == MSO_ANCHOR.MIDDLE:
        ink_y = y + top + max(0.0, (content_height - ink_height) / 2)
    else:
        ink_y = y + top
    return (ink_x, ink_y, ink_width, ink_height)


def is_background(shape: Any, slide_w: float, slide_h: float) -> bool:
    x, y, width, height = box(shape)
    return x <= 0.05 and y <= 0.05 and width >= slide_w * 0.96 and height >= slide_h * 0.96


def is_rule(shape: Any) -> bool:
    if text_of(shape):
        return False
    _, _, width, height = box(shape)
    return shape.shape_type == MSO_SHAPE_TYPE.LINE or (
        max(width, height) >= 0.25 and min(width, height) <= 0.07
    )


def is_meaningful_visual(shape: Any, slide_w: float, slide_h: float) -> bool:
    if shape.shape_type in VISUAL_SHAPE_TYPES:
        return True
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return len(shape.shapes) >= 3
    if shape.shape_type in {
        MSO_SHAPE_TYPE.AUTO_SHAPE,
        MSO_SHAPE_TYPE.FREEFORM,
        MSO_SHAPE_TYPE.LINE,
    }:
        if is_background(shape, slide_w, slide_h):
            return False
        return area(box(shape)) >= 0.08
    return False


def image_ppi(shape: Any) -> float | None:
    display_w = inches(shape.width)
    display_h = inches(shape.height)
    if display_w <= 0 or display_h <= 0:
        return None
    try:
        if shape.image.ext.lower() in {"svg", "emf", "wmf"}:
            return math.inf
        pixel_w, pixel_h = shape.image.size
        return min(pixel_w / display_w, pixel_h / display_h)
    except (AttributeError, ValueError, OSError):
        try:
            relationship_id = shape._pic.blipFill.blip.get(qn("r:embed"))
            if not relationship_id:
                return None
            image_part = shape.part.related_part(relationship_id)
            content_type = image_part.content_type.lower()
            blob = image_part.blob
            if any(
                kind in content_type for kind in ("svg", "emf", "wmf")
            ) or blob.lstrip().startswith((b"<svg", b"<?xml")):
                return math.inf
            with Image.open(io.BytesIO(blob)) as image:
                pixel_w, pixel_h = image.size
            return min(pixel_w / display_w, pixel_h / display_h)
        except (AttributeError, KeyError, ValueError, OSError):
            return None


def load_policy(path: Path | None) -> dict[str, Any]:
    if not path:
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Policy must be a JSON object.")
    return data


def overlaps_allowed(policy: dict[str, Any], slide: int, a: str, b: str) -> bool:
    pairs = policy.get("intentional_overlaps", [])
    wanted = {a, b}
    for item in pairs:
        if (
            isinstance(item, dict)
            and item.get("slide") == slide
            and {item.get("a"), item.get("b")} == wanted
        ):
            return True
    return False


def audit(
    path: Path,
    policy: dict[str, Any] | None = None,
    plan: dict[str, Any] | None = None,
) -> dict[str, Any]:
    policy = policy or {}
    presentation = Presentation(path)
    slide_w = inches(presentation.slide_width)
    slide_h = inches(presentation.slide_height)
    findings: list[Finding] = []
    fingerprints: list[str] = []
    text_only_allowed = set(policy.get("text_only_allowed_slides", []))

    for slide_number, slide in enumerate(presentation.slides, start=1):
        active_shapes = [
            shape
            for shape in slide.shapes
            if area(box(shape)) > 0.0001
            and (text_of(shape) or shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER)
        ]
        text_shapes = [shape for shape in active_shapes if text_of(shape)]
        total_text = sum(len(text_of(shape)) for shape in text_shapes)
        visual_shapes = [
            shape for shape in active_shapes if is_meaningful_visual(shape, slide_w, slide_h)
        ]
        visual_area = sum(
            area(box(shape))
            for shape in visual_shapes
            if not is_background(shape, slide_w, slide_h)
        )
        visual_ratio = min(1.0, visual_area / max(0.1, slide_w * slide_h))
        fingerprint = ":".join(
            f"{kind}={count}"
            for kind, count in sorted(
                Counter(
                    "text"
                    if text_of(shape)
                    else "visual"
                    if is_meaningful_visual(shape, slide_w, slide_h)
                    else "decor"
                    for shape in active_shapes
                ).items()
            )
        )
        fingerprints.append(fingerprint)

        for shape in active_shapes:
            x, y, width, height = box(shape)
            tolerance = 0.015
            if (
                x < -tolerance
                or y < -tolerance
                or x + width > slide_w + tolerance
                or y + height > slide_h + tolerance
            ):
                findings.append(
                    Finding(
                        "error",
                        slide_number,
                        "bounds",
                        (
                            f"Shape bounds ({x:.3f}, {y:.3f}, {width:.3f}, {height:.3f}) "
                            f"exceed slide {slide_w:.3f}x{slide_h:.3f}."
                        ),
                        shape.name,
                    )
                )

            text = text_of(shape)
            if text:
                for pattern in PLACEHOLDER_PATTERNS:
                    if pattern.search(text):
                        findings.append(
                            Finding(
                                "error",
                                slide_number,
                                "placeholder",
                                f"Placeholder-like text remains: {text[:100]!r}",
                                shape.name,
                            )
                        )
                        break
                required, available = estimated_text_height_pt(shape)
                if required > available * 1.12:
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "text-fit",
                            (
                                f"Estimated text height {required:.1f}pt exceeds "
                                f"available {available:.1f}pt."
                            ),
                            shape.name,
                        )
                    )
                sizes = [
                    paragraph_font_size(paragraph)
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ]
                if sizes and min(sizes) < 9:
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "typography",
                            f"Text uses {min(sizes):.1f}pt, below the 9pt absolute minimum.",
                            shape.name,
                        )
                    )
                elif (
                    sizes
                    and min(sizes) < 11
                    and slide_number != 1
                    and not MICROCOPY_NAME.search(shape.name or "")
                ):
                    findings.append(
                        Finding(
                            "warning",
                            slide_number,
                            "typography",
                            f"Text uses {min(sizes):.1f}pt; verify readability at distance.",
                            shape.name,
                        )
                    )
                elif sizes and min(sizes) < 16 and slide_number != 1 and len(text) >= 80:
                    findings.append(
                        Finding(
                            "warning",
                            slide_number,
                            "typography",
                            (
                                f"Long text ({len(text)} chars) uses "
                                f"{min(sizes):.1f}pt; verify readability at distance."
                            ),
                            shape.name,
                        )
                    )

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                ppi = image_ppi(shape)
                if ppi is None:
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "asset",
                            "Picture relation is unreadable.",
                            shape.name,
                        )
                    )
                elif ppi < 110:
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "asset",
                            f"Displayed raster resolution is only {ppi:.0f} PPI.",
                            shape.name,
                        )
                    )
                elif ppi < 140:
                    findings.append(
                        Finding(
                            "warning",
                            slide_number,
                            "asset",
                            f"Displayed raster resolution is {ppi:.0f} PPI.",
                            shape.name,
                        )
                    )
                image_width = inches(shape.width)
                image_height = inches(shape.height)
                if (
                    EVIDENCE_IMAGE_NAME.search(shape.name or "")
                    and "thumbnail" not in (shape.name or "").lower()
                    and (image_width < 2.0 or image_height < 1.0)
                ):
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "asset",
                            (
                                f"Evidence image is too small to be readable "
                                f"({image_width:.2f}x{image_height:.2f} in)."
                            ),
                            shape.name,
                        )
                    )

        for index, first in enumerate(text_shapes):
            first_box = text_ink_box(first)
            for second in text_shapes[index + 1 :]:
                second_box = text_ink_box(second)
                overlap = intersection_area(first_box, second_box)
                if overlap < 0.015:
                    continue
                ratio = overlap / max(0.001, min(area(first_box), area(second_box)))
                if ratio >= 0.02 and not overlaps_allowed(
                    policy, slide_number, first.name, second.name
                ):
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "overlap",
                            (
                                f"Text boxes {first.name!r} and {second.name!r} overlap "
                                f"by {ratio:.1%} of the smaller box."
                            ),
                        )
                    )

        rules = [shape for shape in active_shapes if is_rule(shape)]
        for rule in rules:
            rule_box = box(rule)
            for text_shape in text_shapes:
                ink_box = text_ink_box(text_shape)
                if intersection_area(rule_box, ink_box) >= 0.002 and not overlaps_allowed(
                    policy, slide_number, rule.name, text_shape.name
                ):
                    findings.append(
                        Finding(
                            "error",
                            slide_number,
                            "overlap",
                            (
                                f"Rule {rule.name!r} intersects rendered text in "
                                f"{text_shape.name!r}."
                            ),
                        )
                    )

        containers = [
            shape
            for shape in active_shapes
            if not text_of(shape)
            and not is_rule(shape)
            and not is_background(shape, slide_w, slide_h)
        ]
        for text_shape in text_shapes:
            ink_box = text_ink_box(text_shape)
            center_x = ink_box[0] + ink_box[2] / 2
            center_y = ink_box[1] + ink_box[3] / 2
            candidates = []
            for container in containers:
                container_box = box(container)
                if (
                    container_box[0] <= center_x <= container_box[0] + container_box[2]
                    and container_box[1] <= center_y <= container_box[1] + container_box[3]
                    and area(container_box) >= area(ink_box)
                ):
                    candidates.append((area(container_box), container, container_box))
            if not candidates:
                continue
            _, container, container_box = min(candidates, key=lambda item: item[0])
            tolerance = 0.015
            escaped = (
                ink_box[0] < container_box[0] - tolerance
                or ink_box[1] < container_box[1] - tolerance
                or ink_box[0] + ink_box[2] > container_box[0] + container_box[2] + tolerance
                or ink_box[1] + ink_box[3] > container_box[1] + container_box[3] + tolerance
            )
            if escaped:
                strict = bool(
                    CONTAINMENT_NAME.search(text_shape.name or "")
                    or CONTAINMENT_NAME.search(container.name or "")
                )
                findings.append(
                    Finding(
                        "error" if strict else "warning",
                        slide_number,
                        "containment",
                        (
                            f"Text {text_shape.name!r} escapes its visual container "
                            f"{container.name!r}."
                        ),
                    )
                )

        title_like = slide_number == 1 or (
            total_text <= 110 and len(text_shapes) <= 2 and visual_ratio < 0.1
        )
        if (
            not title_like
            and slide_number not in text_only_allowed
            and total_text >= 70
            and visual_ratio < 0.12
        ):
            findings.append(
                Finding(
                    "error",
                    slide_number,
                    "visual-density",
                    (
                        f"Content slide is effectively text-only "
                        f"(text={total_text} chars, visual area={visual_ratio:.1%})."
                    ),
                )
            )
        if total_text > 850:
            findings.append(
                Finding(
                    "error",
                    slide_number,
                    "density",
                    f"Slide contains {total_text} characters.",
                )
            )
        elif total_text > 520:
            findings.append(
                Finding(
                    "warning",
                    slide_number,
                    "density",
                    f"Slide contains {total_text} characters.",
                )
            )

    for index in range(2, len(fingerprints)):
        if fingerprints[index] == fingerprints[index - 1] == fingerprints[index - 2]:
            findings.append(
                Finding(
                    "warning",
                    index + 1,
                    "rhythm",
                    "Three consecutive slides share the same coarse layout fingerprint.",
                )
            )

    if plan:
        planned_slides = plan.get("slides", [])
        if isinstance(planned_slides, list) and len(planned_slides) != len(presentation.slides):
            findings.append(
                Finding(
                    "error",
                    0,
                    "plan",
                    (
                        f"Plan has {len(planned_slides)} slides but PPTX has "
                        f"{len(presentation.slides)}."
                    ),
                )
            )

    summary = Counter(finding.severity for finding in findings)
    return {
        "path": str(path.resolve()),
        "slide_count": len(presentation.slides),
        "slide_size_inches": {"width": round(slide_w, 4), "height": round(slide_h, 4)},
        "summary": {
            "errors": summary["error"],
            "warnings": summary["warning"],
            "passed": summary["error"] == 0,
        },
        "findings": [asdict(finding) for finding in findings],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit PPTX structure and assets.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--policy", type=Path)
    parser.add_argument("--plan", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--fail-on-warnings", action="store_true")
    args = parser.parse_args()
    try:
        policy = load_policy(args.policy)
        plan = json.loads(args.plan.read_text(encoding="utf-8")) if args.plan else None
        report = audit(args.pptx, policy, plan)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"Quality audit failed: {exc}", file=sys.stderr)
        return 2

    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    if report["summary"]["errors"]:
        return 1
    if args.fail_on_warnings and report["summary"]["warnings"]:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
