import copy
import hashlib
import json
import math
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from scripts.audit_diagram_geometry import audit_file
from scripts.diagram_geometry import audit
from scripts.render_diagram import svg_from_report
from scripts.validate_deck import run_diagram_layout_audits

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"


@pytest.fixture
def layout() -> dict:
    return {
        "width": 600,
        "height": 300,
        "font_family": "Noto Sans CJK JP",
        "font_path": FONT,
        "font_size": 20,
        "line_height": 28,
        "padding": 12,
        "clearance": 8,
        "placement": {
            "width_inches": 600 / 72,
            "height_inches": 300 / 72,
            "min_font_pt": 11,
        },
        "nodes": [
            {
                "id": "a",
                "box": [20, 100, 150, 90],
                "lines": ["入力", "decode"],
                "fill": "#FFFFFF",
                "stroke": "#14213D",
            },
            {
                "id": "b",
                "box": [430, 100, 150, 90],
                "lines": ["出力"],
                "fill": "#FFFFFF",
                "stroke": "#14213D",
            },
        ],
        "edges": [
            {
                "id": "a-b",
                "from": "a",
                "to": "b",
                "points": [[170, 145], [430, 145]],
                "kind": "main",
                "color": "#5E6B78",
                "label": {"text": "YES", "box": [240, 70, 120, 50]},
            },
        ],
    }


def categories(report: dict) -> set[str]:
    assert report["errors"] == len(
        [item for item in report["findings"] if item["severity"] == "error"]
    )
    assert report["passed"] == (report["errors"] == 0)
    json.dumps(report, allow_nan=False)
    return {item["category"] for item in report["findings"]}


def test_valid_layout_has_shared_measurements_without_mutation(layout: dict) -> None:
    before = copy.deepcopy(layout)
    report = audit(layout)
    assert report["passed"], report
    assert layout == before
    assert report["dimensions"]["scale_pt_per_unit"] == pytest.approx(1)
    assert report["fonts"]["minimum_apparent_font_pt"] == pytest.approx(20)
    node = report["measurements"]["nodes"][0]
    assert len(node["lines"]) == 2
    for line in node["lines"]:
        x, y, width, height = line["bbox"]
        assert x >= 32 and x + width <= 158
        assert y >= 112 and y + height <= 178
    assert report == audit(layout)


@pytest.mark.parametrize(
    "field,value",
    [
        ("width", math.nan),
        ("height", math.inf),
        ("font_size", -1),
        ("font_size", True),
        ("line_height", 0),
        ("padding", -1),
        ("clearance", -1),
        ("nodes", []),
        ("nodes", {}),
        ("edges", {}),
        ("font_family", ""),
        ("font_path", "/no/such/font.ttf"),
        ("placement", {}),
        ("width", "600"),
        ("width", 10**1000),
    ],
)
def test_bad_root_values_are_reports(layout: dict, field: str, value: object) -> None:
    layout[field] = value
    assert categories(audit(layout))


@pytest.mark.parametrize("value", [None, [], "", 12])
def test_non_object_input_is_rejected(value: object) -> None:
    assert "input" in categories(audit(value))


@pytest.mark.parametrize(
    "change",
    [
        {"box": [1, 2, 3]},
        {"box": [0, 0, 0, 20]},
        {"box": [0, 0, math.nan, 20]},
        {"box": "wrong"},
        {"lines": "text"},
        {"lines": []},
        {"lines": [42]},
        {"shape": "diamond"},
        {"id": ""},
        {"font_size": 0},
        {"fill": 'url("evil")'},
        {"lines": ["bad\x00text"]},
    ],
)
def test_bad_nodes_are_reports(layout: dict, change: dict) -> None:
    layout["nodes"][0].update(change)
    assert categories(audit(layout))


@pytest.mark.parametrize(
    "change",
    [
        {"points": []},
        {"points": [[1, 2]]},
        {"points": [[1, 2, 3], [4, 5]]},
        {"points": [[1, math.inf], [4, 5]]},
        {"points": [[170, 145], [170, 145], [430, 145]]},
        {"points": [[170, 145], [430, 146]]},
        {"kind": "curve"},
        {"from": "unknown"},
        {"to": None},
        {"id": ""},
        {"label": "YES"},
        {"label": {}},
        {"label": {"text": "YES", "box": [1, 2, -1, 4]}},
    ],
)
def test_bad_edges_are_reports(layout: dict, change: dict) -> None:
    layout["edges"][0].update(change)
    assert categories(audit(layout))


def test_duplicate_ids_and_missing_required_fields(layout: dict) -> None:
    layout["nodes"][1]["id"] = "a"
    assert "input" in categories(audit(layout))
    layout["nodes"][1]["id"] = "b"
    layout["edges"].append(copy.deepcopy(layout["edges"][0]))
    assert "input" in categories(audit(layout))
    del layout["font_size"]
    assert "input" in categories(audit(layout))


@pytest.mark.parametrize(
    "text", ["日本語の長い文字列が枠からはみ出します", "very_long_code_identifier_without_spaces"]
)
def test_actual_text_bbox_overflow(layout: dict, text: str) -> None:
    layout["nodes"][0]["lines"] = [text]
    assert "text-containment" in categories(audit(layout))


def test_multiline_overlap_and_label_insets(layout: dict) -> None:
    layout["line_height"] = 5
    assert "text-overlap" in categories(audit(layout))
    layout["line_height"] = 28
    layout["edges"][0]["label"]["box"] = [240, 70, 40, 20]
    assert "text-containment" in categories(audit(layout))


def test_declared_contain_scale_and_overrides(layout: dict) -> None:
    layout["placement"]["width_inches"] = 12
    layout["placement"]["height_inches"] = 2
    report = audit(layout)
    assert report["dimensions"]["expected_width_inches"] == pytest.approx(4)
    assert report["dimensions"]["expected_height_inches"] == pytest.approx(2)
    assert "minimum-font-size" in categories(report)
    layout["placement"]["height_inches"] = 5
    layout["nodes"][0]["font_size"] = 9
    assert "minimum-font-size" in categories(audit(layout))
    layout["nodes"][0]["font_size"] = 20
    layout["edges"][0]["label"]["font_size"] = 8
    assert "minimum-font-size" in categories(audit(layout))


def test_agreed_canvas_font_threshold_is_not_rounded_up(layout: dict) -> None:
    layout.update(width=1440, height=570, font_size=22)
    layout["placement"].update(width_inches=11.95, height_inches=3.95, min_font_pt=11)
    report = audit(layout)
    assert report["fonts"]["minimum_apparent_font_pt"] == pytest.approx(10.976842105263158)
    assert "minimum-font-size" in categories(report)
    layout["font_size"] = 23
    assert audit(layout)["passed"], audit(layout)


def test_bounds_include_stroke_labels_and_arrowheads(layout: dict) -> None:
    layout["nodes"][0]["box"][0] = 0
    assert "canvas-bounds" in categories(audit(layout))
    layout["nodes"][0]["box"][0] = 20
    layout["edges"][0]["label"]["box"][0] = 590
    assert "canvas-bounds" in categories(audit(layout))
    layout["edges"][0].pop("label")
    layout["edges"][0]["points"] = [
        [95, 100],
        [95, 2],
        [505, 2],
        [505, 100],
    ]
    assert audit(layout)["passed"]
    layout["nodes"].append({"id": "j", "shape": "junction", "box": [400, 1, 4, 4], "lines": []})
    layout["edges"][0].update(to="j", points=[[95, 100], [95, 3], [400, 3]])
    assert "canvas-bounds" in categories(audit(layout))


def test_node_clearance_includes_strokes(layout: dict) -> None:
    layout["nodes"][1]["box"] = [179, 100, 150, 90]
    layout["edges"] = []
    assert "node-clearance" in categories(audit(layout))
    layout["nodes"][1]["box"][0] = 180
    assert audit(layout)["passed"]


def test_return_kind_is_not_a_geometry_exemption(layout: dict) -> None:
    layout["nodes"].append({"id": "middle", "box": [240, 125, 120, 60], "lines": ["middle"]})
    layout["edges"][0].update(kind="return", route="outside")
    assert "edge-node" in categories(audit(layout))


def test_edge_respects_own_and_unrelated_label_clearance(layout: dict) -> None:
    layout["edges"][0]["label"]["box"] = [240, 100, 120, 50]
    assert "edge-label" in categories(audit(layout))
    layout["edges"][0]["label"]["box"] = [240, 87, 120, 50]
    assert "edge-label" in categories(audit(layout))
    layout["edges"][0]["label"]["box"] = [240, 86, 120, 50]
    assert audit(layout)["passed"]
    layout["nodes"].append({"id": "c", "shape": "junction", "box": [398, 30, 4, 4], "lines": []})
    layout["edges"].append(
        {
            "id": "c-b",
            "from": "c",
            "to": "b",
            "kind": "branch",
            "points": [[400, 34], [400, 145], [430, 145]],
            "label": {"text": "NO", "box": [240, 110, 100, 50]},
        }
    )
    assert "edge-label" in categories(audit(layout))


def test_label_node_and_label_label_clearance(layout: dict) -> None:
    layout["edges"][0]["label"]["box"] = [30, 40, 100, 55]
    assert "label-node" in categories(audit(layout))
    layout["edges"].append(
        {
            "id": "b-a",
            "from": "b",
            "to": "a",
            "kind": "return",
            "points": [[505, 100], [505, 30], [95, 30], [95, 100]],
            "label": {"text": "NO", "box": [250, 72, 120, 50]},
        }
    )
    layout["edges"][0]["label"]["box"] = [240, 70, 120, 50]
    assert "label-label" in categories(audit(layout))


def test_endpoint_attachment_and_arrow_length(layout: dict) -> None:
    layout["edges"][0]["points"][0][0] = 171
    assert "endpoint-attachment" in categories(audit(layout))
    layout["edges"][0]["points"] = [[170, 145], [425, 145], [430, 145]]
    assert "arrowhead-length" in categories(audit(layout))


def test_endpoint_cannot_enter_or_traverse_its_own_node(layout: dict) -> None:
    layout["edges"][0]["points"] = [
        [20, 145],
        [300, 145],
        [300, 220],
        [505, 220],
        [505, 190],
    ]
    assert "edge-node" in categories(audit(layout))
    layout["edges"][0]["points"] = [
        [170, 145],
        [200, 145],
        [200, 120],
        [50, 120],
        [50, 220],
        [505, 220],
        [505, 190],
    ]
    assert "edge-node" in categories(audit(layout))


def test_arrowhead_geometry_not_just_centerline(layout: dict) -> None:
    layout["padding"] = 0
    layout["clearance"] = 0
    layout["edges"][0].pop("label")
    layout["nodes"].append(
        {"id": "near-arrow", "shape": "junction", "box": [420, 146, 3, 3], "lines": []}
    )
    assert "edge-node" in categories(audit(layout))


def test_crossing_and_collinear_overlap_are_detected(layout: dict) -> None:
    layout["nodes"] += [
        {"id": "c", "box": [220, 10, 150, 50], "lines": ["C"]},
        {"id": "d", "box": [220, 230, 150, 50], "lines": ["D"]},
    ]
    layout["edges"][0].pop("label")
    layout["edges"].append(
        {"id": "c-d", "from": "c", "to": "d", "kind": "main", "points": [[295, 60], [295, 230]]}
    )
    assert "edge-crossing" in categories(audit(layout))
    layout["edges"][1]["points"] = [
        [295, 60],
        [295, 145],
        [400, 145],
        [400, 220],
        [295, 220],
        [295, 230],
    ]
    assert "edge-overlap" in categories(audit(layout))


def test_shared_declared_node_does_not_excuse_overlap(layout: dict) -> None:
    edge = copy.deepcopy(layout["edges"][0])
    edge["id"] = "second"
    layout["edges"].append(edge)
    assert "edge-overlap" in categories(audit(layout))


def test_self_crossings_are_detected(layout: dict) -> None:
    layout["edges"][0].pop("label")
    layout["edges"][0]["points"] = [
        [170, 145],
        [300, 145],
        [300, 220],
        [250, 220],
        [250, 70],
        [380, 70],
        [380, 145],
        [430, 145],
    ]
    assert "edge-crossing" in categories(audit(layout))


def test_junction_intentionally_merges_separate_ports(layout: dict) -> None:
    layout["nodes"] = [
        {"id": "a", "box": [20, 30, 150, 60], "lines": ["A"]},
        {"id": "b", "box": [20, 210, 150, 60], "lines": ["B"]},
        {"id": "j", "shape": "junction", "box": [290, 142, 16, 16], "lines": []},
        {"id": "c", "box": [430, 120, 150, 60], "lines": ["C"]},
    ]
    layout["edges"] = [
        {
            "id": "a-j",
            "from": "a",
            "to": "j",
            "kind": "branch",
            "points": [[170, 60], [298, 60], [298, 142]],
        },
        {
            "id": "b-j",
            "from": "b",
            "to": "j",
            "kind": "branch",
            "points": [[170, 240], [298, 240], [298, 158]],
        },
        {"id": "j-c", "from": "j", "to": "c", "kind": "main", "points": [[306, 150], [430, 150]]},
    ]
    assert audit(layout)["passed"], audit(layout)


def write_layout(tmp_path: Path, layout: dict) -> Path:
    path = tmp_path / "flow.layout.json"
    path.write_text(json.dumps(layout, ensure_ascii=False), encoding="utf-8")
    return path


def test_renderer_is_deterministic_escaped_and_uses_audited_geometry(
    tmp_path: Path,
    layout: dict,
) -> None:
    layout["nodes"][0]["lines"] = ['A<&"B']
    path = write_layout(tmp_path, layout)
    output = tmp_path / "flow.svg"
    report_path = tmp_path / "report.json"
    command = [
        sys.executable,
        str(SCRIPTS / "render_diagram.py"),
        str(path),
        "--output",
        str(output),
        "--report",
        str(report_path),
    ]
    first = subprocess.run(command, capture_output=True, text=True)
    assert first.returncode == 0, first.stdout + first.stderr
    content = output.read_bytes()
    assert subprocess.run(command, capture_output=True).returncode == 0
    assert content == output.read_bytes()
    root = ET.fromstring(content)
    ns = {"s": "http://www.w3.org/2000/svg"}
    report = json.loads(report_path.read_text())
    texts = root.findall(".//s:text", ns)
    assert texts[0].text == 'A<&"B'
    line = report["measurements"]["nodes"][0]["lines"][0]
    assert float(texts[0].attrib["x"]) == pytest.approx(line["x"], abs=1e-6)
    assert float(texts[0].attrib["y"]) == pytest.approx(line["y"], abs=1e-6)
    assert root.findall(".//s:polygon", ns)
    assert not root.findall(".//s:marker", ns)
    assert report["sha256"]["layout"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_failed_render_preserves_output_and_writes_failed_report(
    tmp_path: Path,
    layout: dict,
) -> None:
    layout["nodes"][0]["lines"] = ["x" * 300]
    path = write_layout(tmp_path, layout)
    output = tmp_path / "flow.svg"
    output.write_bytes(b"previous-good-output")
    report_path = tmp_path / "report.json"
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPTS / "render_diagram.py"),
            str(path),
            "--output",
            str(output),
            "--report",
            str(report_path),
        ],
        capture_output=True,
    )
    assert result.returncode != 0
    assert output.read_bytes() == b"previous-good-output"
    assert not json.loads(report_path.read_text())["passed"]


def make_picture_deck(
    tmp_path: Path, *, width: float = 600 / 72, height: float = 300 / 72
) -> tuple[Path, Path]:
    image = tmp_path / "flow.png"
    Image.new("RGB", (1200, 600), "white").save(image)
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(14)
    presentation.slide_height = Inches(8)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = slide.shapes.add_picture(
        str(image), Inches(1), Inches(1), width=Inches(width), height=Inches(height)
    )
    picture.name = "main-graph-external"
    presentation.save(deck)
    return deck, image


def write_provenance(layout_path: Path, image_path: Path) -> Path:
    svg = image_path.with_suffix(".svg")
    svg.write_text(svg_from_report(audit(json.loads(layout_path.read_bytes()))), encoding="utf-8")
    provenance = image_path.with_suffix(".provenance.json")
    provenance.write_text(
        json.dumps(
            {
                "layout_sha256": hashlib.sha256(layout_path.read_bytes()).hexdigest(),
                "svg_sha256": hashlib.sha256(svg.read_bytes()).hexdigest(),
                "image_sha256": hashlib.sha256(image_path.read_bytes()).hexdigest(),
                "renderer_sha256": hashlib.sha256(
                    (SCRIPTS / "diagram_geometry.py").read_bytes()
                ).hexdigest(),
            }
        )
    )
    return provenance


def test_integrated_contain_binding_and_actual_font_override(
    tmp_path: Path,
    layout: dict,
) -> None:
    layout["placement"].update(width_inches=12, height_inches=300 / 72)
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    report = audit_file(
        path,
        image_path=image,
        pptx_path=deck,
        slide=1,
        shape_name="main-graph-external",
        provenance_path=provenance,
    )
    assert report["passed"], report
    assert report["sha256"]["image"] == hashlib.sha256(image.read_bytes()).hexdigest()
    assert report["embedded_image_sha256"] == report["sha256"]["image"]
    assert report["actual_placement"]["minimum_apparent_font_pt"] == pytest.approx(20)
    layout["nodes"][0]["font_size"] = 9
    write_layout(tmp_path, layout)
    assert "minimum-font-size" in categories(
        audit_file(
            path,
            image_path=image,
            pptx_path=deck,
            slide=1,
            shape_name="main-graph-external",
        )
    )


def test_integrated_placement_mismatch_and_stale_image(
    tmp_path: Path,
    layout: dict,
) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path, width=4, height=2)
    report = audit_file(
        path, image_path=image, pptx_path=deck, slide=1, shape_name="main-graph-external"
    )
    assert "image-placement" in categories(report)
    assert "minimum-font-size" in categories(report)
    Image.new("RGB", (1200, 600), "red").save(image)
    assert "image-content" in categories(
        audit_file(
            path,
            image_path=image,
            pptx_path=deck,
            slide=1,
            shape_name="main-graph-external",
        )
    )


@pytest.mark.parametrize(
    "slide,shape", [(2, "main-graph-external"), (1, "missing"), (0, "main-graph-external")]
)
def test_integrated_missing_picture_is_not_skipped(
    tmp_path: Path,
    layout: dict,
    slide: int,
    shape: str,
) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    assert categories(
        audit_file(path, image_path=image, pptx_path=deck, slide=slide, shape_name=shape)
    )


def test_integrated_font_policy(tmp_path: Path, layout: dict) -> None:
    path = write_layout(tmp_path, layout)
    policy = {"enabled": True, "primary_font": "Arial", "mono_font": "Consolas"}
    assert "font-family" in categories(audit_file(path, typography_policy=policy))
    policy["primary_font"] = layout["font_family"]
    assert audit_file(path, typography_policy=policy)["passed"]


def test_validate_deck_runs_geometry_cli_and_captures_hashes(
    tmp_path: Path,
    layout: dict,
) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    plan_path = tmp_path / "plan.json"
    plan = {
        "meta": {
            "diagram_layouts": [
                {
                    "path": path.name,
                    "image": image.name,
                    "provenance": provenance.name,
                    "slide": 1,
                    "shape": "main-graph-external",
                },
            ]
        }
    }
    plan_path.write_text(json.dumps(plan))
    output = tmp_path / "qa"
    output.mkdir()
    reports = run_diagram_layout_audits(plan, plan_path, deck, output, {})
    assert len(reports) == 1
    entry = reports[0]
    assert Path(entry["report"]).is_file()
    assert entry["sha256"]["layout"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert entry["sha256"]["image"] == hashlib.sha256(image.read_bytes()).hexdigest()
    assert (
        entry["sha256"]["report"] == hashlib.sha256(Path(entry["report"]).read_bytes()).hexdigest()
    )
    for key in ("provenance", "svg", "renderer"):
        assert Path(entry[key]).is_file()
        assert entry["sha256"][key] == hashlib.sha256(Path(entry[key]).read_bytes()).hexdigest()


@pytest.mark.parametrize(
    "field,value",
    [
        ("line_height", 1e308),
        ("width", 1e-300),
        ("font_size", 1e308),
    ],
)
def test_unrepresentable_geometry_returns_json_safe_errors(
    layout: dict, field: str, value: float
) -> None:
    layout[field] = value
    layout["nodes"][0]["lines"] = ["A", "B", "C"]
    assert categories(audit(layout))


def test_unrepresentable_boxes_and_segments_do_not_raise(layout: dict) -> None:
    layout["nodes"][0]["box"] = [1e308, 100, 1, 90]
    assert categories(audit(layout))
    layout["nodes"][0]["box"] = [20, 100, 1e-300, 90]
    assert categories(audit(layout))
    layout["nodes"][0]["box"] = [20, 100, 150, 90]
    layout["edges"][0]["points"] = [[1e-300, 0], [2e-300, 0]]
    assert categories(audit(layout))


def test_stroke_overlap_between_parallel_edges_is_detected(layout: dict) -> None:
    layout["padding"] = 0
    layout["clearance"] = 0
    layout["edges"][0].pop("label")
    layout["edges"].append(
        {
            "id": "return",
            "from": "b",
            "to": "a",
            "kind": "return",
            "points": [[505, 100], [505, 60], [95, 60], [95, 100]],
        }
    )
    layout["edges"][0]["points"] = [
        [170, 145],
        [200, 145],
        [200, 61.5],
        [400, 61.5],
        [400, 145],
        [430, 145],
    ]
    assert "edge-overlap" in categories(audit(layout))


def test_arrow_stroke_contact_is_detected(layout: dict) -> None:
    layout["edges"][0].pop("label")
    layout["nodes"] += [
        {"id": "c", "box": [215, 15, 120, 50], "lines": ["C"]},
        {"id": "j", "shape": "junction", "box": [395, 130, 4, 4], "lines": []},
    ]
    layout["edges"].append(
        {
            "id": "c-j",
            "from": "c",
            "to": "j",
            "kind": "branch",
            "points": [[335, 40], [397, 40], [397, 130]],
        }
    )
    assert audit(layout)["passed"]
    # A triangle reaches y=136 at its base while the unrelated line is y=136.5.
    layout["edges"][0]["points"] = [
        [170, 145],
        [200, 145],
        [200, 136.5],
        [410, 136.5],
        [410, 145],
        [430, 145],
    ]
    layout["edges"][1]["points"] = [[335, 40], [365, 40], [365, 132], [395, 132]]
    assert "edge-crossing" in categories(audit(layout))


def test_shared_endpoint_only_is_allowed(layout: dict) -> None:
    layout["edges"][0].pop("label")
    layout["nodes"] = [
        {"id": "j", "shape": "junction", "box": [290, 142, 16, 16], "lines": []},
        {"id": "a", "box": [20, 112, 150, 60], "lines": ["A"]},
        {"id": "b", "box": [222, 210, 150, 60], "lines": ["B"]},
    ]
    layout["edges"] = [
        {"id": "j-a", "from": "j", "to": "a", "kind": "main", "points": [[290, 158], [170, 158]]},
        {"id": "j-b", "from": "j", "to": "b", "kind": "branch", "points": [[290, 158], [290, 210]]},
    ]
    assert audit(layout)["passed"], audit(layout)


def test_actual_font_must_pass_even_within_size_tolerance(tmp_path: Path, layout: dict) -> None:
    layout["placement"]["min_font_pt"] = 20
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path, width=600 / 72 - 0.01, height=300 / 72 - 0.005)
    report = audit_file(
        path, image_path=image, pptx_path=deck, slide=1, shape_name="main-graph-external"
    )
    assert "image-placement" not in categories(report)
    assert "minimum-font-size" in categories(report)


@pytest.mark.parametrize("change", ["crop", "rotate", "flip", "duplicate", "wrong-aspect"])
def test_transformed_or_ambiguous_images_are_rejected(
    tmp_path: Path, layout: dict, change: str
) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    presentation = Presentation(deck)
    picture = presentation.slides[0].shapes[0]
    if change == "crop":
        picture.crop_left = 0.1
    elif change == "rotate":
        picture.rotation = 180
    elif change == "flip":
        picture._element.spPr.xfrm.set("flipH", "1")
    elif change == "duplicate":
        other = presentation.slides[0].shapes.add_picture(str(image), 0, 0)
        other.name = picture.name
    else:
        Image.new("RGB", (100, 100), "white").save(image)
    presentation.save(deck)
    assert categories(
        audit_file(
            path, image_path=image, pptx_path=deck, slide=1, shape_name="main-graph-external"
        )
    )


def test_explicit_typography_range_is_preserved(tmp_path: Path, layout: dict) -> None:
    path = write_layout(tmp_path, layout)
    policy = {
        "enabled": True,
        "primary_font": layout["font_family"],
        "generated_diagram_apparent_font_range": [11, 18],
    }
    assert "generated-diagram-apparent-font-size" in categories(
        audit_file(
            path,
            typography_policy=policy,
        )
    )


@pytest.mark.parametrize("text", ["null", "{", '{"width": NaN}', '["not a layout"]'])
def test_both_clis_report_invalid_json_without_touching_svg(tmp_path: Path, text: str) -> None:
    path = tmp_path / "bad.layout.json"
    path.write_text(text)
    for script in ("render_diagram.py", "audit_diagram_geometry.py"):
        report_path = tmp_path / "bad-report.json"
        command = [sys.executable, str(SCRIPTS / script), str(path)]
        if script == "render_diagram.py":
            output = tmp_path / "old.svg"
            output.write_bytes(b"old")
            command += ["--output", str(output), "--report", str(report_path)]
        else:
            command += ["--output", str(report_path)]
        result = subprocess.run(command, capture_output=True)
        assert result.returncode != 0
        assert not json.loads(report_path.read_text())["passed"]
        assert output.read_bytes() == b"old"


def test_freshness_includes_geometry_artifacts(tmp_path: Path, layout: dict) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    plan_path = tmp_path / "plan.json"
    plan = {
        "meta": {
            "diagram_layouts": [
                {
                    "path": path.name,
                    "image": image.name,
                    "provenance": provenance.name,
                    "slide": 1,
                    "shape": "main-graph-external",
                },
            ]
        }
    }
    plan_path.write_text(json.dumps(plan))
    records = run_diagram_layout_audits(plan, plan_path, deck, tmp_path, {})
    renderer_copy = tmp_path / "renderer-copy.py"
    renderer_copy.write_bytes((SCRIPTS / "diagram_geometry.py").read_bytes())
    records[0]["renderer"] = str(renderer_copy)
    summary = tmp_path / "qa-summary.json"
    summary.write_text(json.dumps({"diagram_geometry_reports": records}))
    command = [sys.executable, str(SCRIPTS / "verify_qa_freshness.py"), str(summary)]
    assert subprocess.run(command, capture_output=True).returncode == 0
    for key in ("layout", "image", "report", "provenance", "svg", "renderer"):
        artifact = Path(records[0][key])
        original = artifact.read_bytes()
        artifact.write_bytes(original + b"\n")
        result = subprocess.run(command, capture_output=True)
        assert result.returncode != 0, key
        artifact.write_bytes(original)
    path.unlink()
    assert subprocess.run(command, capture_output=True).returncode != 0


def test_image_audit_requires_provenance_but_geometry_only_does_not(
    tmp_path: Path,
    layout: dict,
) -> None:
    path = write_layout(tmp_path, layout)
    _, image = make_picture_deck(tmp_path)
    assert audit_file(path)["passed"]
    assert "provenance" in categories(audit_file(path, image_path=image))
    provenance = write_provenance(path, image)
    assert audit_file(path, image_path=image, provenance_path=provenance)["passed"]
    assert "provenance" in categories(audit_file(path, provenance_path=provenance))


@pytest.mark.parametrize("key", ["layout", "svg", "image", "renderer"])
def test_provenance_hash_mismatch_is_rejected(tmp_path: Path, layout: dict, key: str) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    manifest = json.loads(provenance.read_bytes())
    manifest[f"{key}_sha256"] = "0" * 64
    provenance.write_text(json.dumps(manifest))
    report = audit_file(
        path,
        image_path=image,
        pptx_path=deck,
        slide=1,
        shape_name="main-graph-external",
        provenance_path=provenance,
    )
    assert "provenance" in categories(report)
    assert any(key in finding["message"] for finding in report["findings"])


@pytest.mark.parametrize("contents", ["null", "[]", "{}", "{", '{"layout_sha256": 7}'])
def test_malformed_provenance_reports_error(tmp_path: Path, layout: dict, contents: str) -> None:
    path = write_layout(tmp_path, layout)
    _, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    provenance.write_text(contents)
    assert "provenance" in categories(
        audit_file(path, image_path=image, provenance_path=provenance)
    )


def test_json_only_change_rejects_stale_image_even_when_geometry_passes(
    tmp_path: Path, layout: dict
) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    layout["nodes"][0]["lines"] = ["Changed"]
    write_layout(tmp_path, layout)
    assert audit(layout)["passed"]
    report = audit_file(
        path,
        image_path=image,
        pptx_path=deck,
        slide=1,
        shape_name="main-graph-external",
        provenance_path=provenance,
    )
    assert "provenance" in categories(report)
    # Merely republishing the new JSON hash must not validate the old SVG.
    manifest = json.loads(provenance.read_bytes())
    manifest["layout_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    provenance.write_text(json.dumps(manifest))
    assert "provenance" in categories(
        audit_file(path, image_path=image, provenance_path=provenance)
    )


def test_svg_hash_republication_does_not_hide_stale_svg(tmp_path: Path, layout: dict) -> None:
    path = write_layout(tmp_path, layout)
    _, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    svg = image.with_suffix(".svg")
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg"/>')
    manifest = json.loads(provenance.read_bytes())
    manifest["svg_sha256"] = hashlib.sha256(svg.read_bytes()).hexdigest()
    provenance.write_text(json.dumps(manifest))
    report = audit_file(path, image_path=image, provenance_path=provenance)
    assert "provenance" in categories(report)
    assert any("current layout" in finding["message"] for finding in report["findings"])


@pytest.mark.parametrize("missing", ["svg", "provenance"])
def test_missing_provenance_artifact_is_rejected(
    tmp_path: Path, layout: dict, missing: str
) -> None:
    path = write_layout(tmp_path, layout)
    _, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    (image.with_suffix(".svg") if missing == "svg" else provenance).unlink()
    assert "provenance" in categories(
        audit_file(path, image_path=image, provenance_path=provenance)
    )


def test_declared_layout_cannot_omit_provenance(tmp_path: Path, layout: dict) -> None:
    path = write_layout(tmp_path, layout)
    deck, image = make_picture_deck(tmp_path)
    plan_path = tmp_path / "plan.json"
    plan = {
        "meta": {
            "diagram_layouts": [
                {
                    "path": path.name,
                    "image": image.name,
                    "slide": 1,
                    "shape": "main-graph-external",
                },
            ]
        }
    }
    plan_path.write_text(json.dumps(plan))
    with pytest.raises(ValueError, match="provenance"):
        run_diagram_layout_audits(plan, plan_path, deck, tmp_path, {})


def test_provenance_cli_success_and_input_overwrite_protection(
    tmp_path: Path, layout: dict
) -> None:
    path = write_layout(tmp_path, layout)
    _, image = make_picture_deck(tmp_path)
    provenance = write_provenance(path, image)
    report_path = tmp_path / "report.json"
    command = [
        sys.executable,
        str(SCRIPTS / "audit_diagram_geometry.py"),
        str(path),
        "--image",
        str(image),
        "--provenance",
        str(provenance),
    ]
    result = subprocess.run([*command, "--output", str(report_path)], capture_output=True)
    assert result.returncode == 0, result.stderr
    assert json.loads(report_path.read_bytes())["passed"]
    for target in (provenance, image.with_suffix(".svg")):
        original = target.read_bytes()
        assert (
            subprocess.run([*command, "--output", str(target)], capture_output=True).returncode != 0
        )
        assert target.read_bytes() == original
