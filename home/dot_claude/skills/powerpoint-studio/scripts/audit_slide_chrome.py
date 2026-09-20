#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

PAGE_NUMBER_RE = re.compile(r"^\s*\d+\s*(?:/\s*\d+)?\s*$")


def load_json(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("The plan root must be a JSON object.")
    return data


def shape_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text.strip()


def audit(pptx_path: Path, plan: dict[str, Any]) -> dict[str, Any]:
    meta = plan.get("meta")
    policy = meta.get("chrome_policy", {}) if isinstance(meta, dict) else {}
    enabled = isinstance(policy, dict) and policy.get("enabled") is True
    if not enabled:
        return {
            "enabled": False,
            "passed": True,
            "errors": 0,
            "findings": [],
        }

    header_allowed = policy.get("header") == "user-requested"
    footer_allowed = policy.get("footer") == "user-requested"
    page_numbers_required = policy.get("page_numbers", "required") == "required"
    top_zone = Inches(float(policy.get("header_zone_inches", 0.55)))
    bottom_zone = Inches(float(policy.get("footer_zone_inches", 0.45)))
    presentation = Presentation(pptx_path)
    bottom_start = presentation.slide_height - bottom_zone
    findings: list[dict[str, Any]] = []

    def add(slide: int, category: str, message: str, text: str = "") -> None:
        findings.append(
            {
                "severity": "error",
                "slide": slide,
                "category": category,
                "message": message,
                "text": text,
            }
        )

    for number, slide in enumerate(presentation.slides, start=1):
        page_numbers = 0
        for shape in slide.shapes:
            text = shape_text(shape)
            if not text:
                continue
            is_page_number = bool(PAGE_NUMBER_RE.fullmatch(text))
            if is_page_number and shape.top >= bottom_start:
                page_numbers += 1
                continue
            if shape.top + shape.height <= top_zone and not header_allowed:
                add(
                    number,
                    "unexpected-header",
                    "Default decks must not contain a small header band or header text.",
                    text,
                )
            if shape.top >= bottom_start and not footer_allowed:
                add(
                    number,
                    "unexpected-footer",
                    "Default decks must not contain footer text other than the page number.",
                    text,
                )
        if page_numbers_required and page_numbers != 1:
            add(
                number,
                "missing-page-number",
                f"Expected exactly one page number in the bottom zone; found {page_numbers}.",
            )

    errors = len(findings)
    return {
        "enabled": True,
        "passed": errors == 0,
        "errors": errors,
        "policy": {
            "header": policy.get("header", "none"),
            "footer": policy.get("footer", "none"),
            "page_numbers": policy.get("page_numbers", "required"),
        },
        "findings": findings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit decorative headers, footers, and required page numbers."
    )
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    try:
        report = audit(args.pptx, load_json(args.plan))
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        report = {
            "enabled": True,
            "passed": False,
            "errors": 1,
            "findings": [
                {
                    "severity": "error",
                    "slide": 0,
                    "category": "input",
                    "message": str(exc),
                }
            ],
        }
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
