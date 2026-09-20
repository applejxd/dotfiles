from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from scripts.template_fidelity import compare


def make_deck(path: Path, width: float = 10.0) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(path)


def test_identical_template_passes(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    make_deck(template)
    report = compare(template, template)
    assert report["passed"]
    assert report["warnings"] == 0


def test_changed_slide_size_fails(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "output.pptx"
    make_deck(template, 10.0)
    make_deck(output, 13.333)
    report = compare(template, output)
    assert not report["passed"]
    assert any(item["category"] == "slide-size" for item in report["findings"])
