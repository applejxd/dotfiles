#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation

CJK = re.compile(r"[\u3040-\u30ff\u3400-\u9fff]")
LATIN_WORD = re.compile(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*")
PAUSE = re.compile(r"[。、，．！？!?：:；;]")


def notes_text(slide: Any) -> str:
    try:
        return slide.notes_slide.notes_text_frame.text.strip()
    except (AttributeError, KeyError):
        return ""


def speech_diagnostic_seconds(text: str) -> float:
    cjk_characters = len(CJK.findall(text))
    latin_words = len(LATIN_WORD.findall(text))
    pauses = len(PAUSE.findall(text))
    return cjk_characters / 280 * 60 + latin_words / 130 * 60 + pauses * 0.12


def validate(pptx_path: Path, plan_path: Path) -> dict[str, Any]:
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    planned_slides = plan.get("slides", [])
    presentation = Presentation(pptx_path)
    findings: list[dict[str, Any]] = []
    timing_total = 0.0
    optional_cut_total = 0.0
    slide_reports = []

    if len(planned_slides) != len(presentation.slides):
        findings.append(
            {
                "severity": "error",
                "slide": 0,
                "category": "timing",
                "message": "Plan and PPTX slide counts differ.",
            }
        )

    for index, (planned, slide) in enumerate(
        zip(planned_slides, presentation.slides, strict=False), start=1
    ):
        timing = planned.get("timing_seconds")
        if timing is None:
            continue
        timing = float(timing)
        optional_cut = float(planned.get("optional_cut_seconds", 0))
        timing_total += timing
        optional_cut_total += optional_cut
        notes = notes_text(slide)
        diagnostic = speech_diagnostic_seconds(notes)
        if len(re.sub(r"\s+", "", notes)) < 40:
            findings.append(
                {
                    "severity": "error",
                    "slide": index,
                    "category": "speaker-notes",
                    "message": "Speaker notes are missing or too short to guide delivery.",
                }
            )
        elif diagnostic < timing * 0.3:
            findings.append(
                {
                    "severity": "warning",
                    "slide": index,
                    "category": "speaker-notes",
                    "message": (
                        f"Notes diagnostic is {diagnostic:.0f}s for a {timing:.0f}s "
                        "slot; add examples, pointing order, pauses, or audience interaction."
                    ),
                }
            )
        elif diagnostic > timing * 1.35:
            findings.append(
                {
                    "severity": "warning",
                    "slide": index,
                    "category": "speaker-notes",
                    "message": (
                        f"Notes diagnostic is {diagnostic:.0f}s for a {timing:.0f}s "
                        "slot; identify optional cuts."
                    ),
                }
            )
        slide_reports.append(
            {
                "slide": index,
                "timing_seconds": timing,
                "optional_cut_seconds": optional_cut,
                "notes_characters": len(re.sub(r"\s+", "", notes)),
                "speech_diagnostic_seconds": round(diagnostic, 1),
            }
        )

    meta = plan.get("meta", {})
    duration = meta.get("duration_minutes")
    if (
        timing_total
        and isinstance(duration, (int, float))
        and abs(timing_total - float(duration) * 60) > 5
    ):
        findings.append(
            {
                "severity": "error",
                "slide": 0,
                "category": "timing",
                "message": "Slide timing total does not match meta.duration_minutes.",
            }
        )
    session_duration = meta.get("session_duration_minutes")
    buffer_minutes = meta.get("buffer_minutes")
    if timing_total >= 20 * 60 and session_duration is None:
        findings.append(
            {
                "severity": "warning",
                "slide": 0,
                "category": "timing",
                "message": "Long-form talk has no meta.session_duration_minutes.",
            }
        )
    if timing_total >= 20 * 60 and not meta.get("checkpoints"):
        findings.append(
            {
                "severity": "warning",
                "slide": 0,
                "category": "timing",
                "message": "Long-form talk has no checkpoint plan in meta.checkpoints.",
            }
        )

    errors = sum(item["severity"] == "error" for item in findings)
    return {
        "pptx": str(pptx_path.resolve()),
        "plan": str(plan_path.resolve()),
        "passed": errors == 0,
        "errors": errors,
        "warnings": sum(item["severity"] == "warning" for item in findings),
        "timing_total_seconds": timing_total,
        "optional_cut_total_seconds": optional_cut_total,
        "session_duration_minutes": session_duration,
        "buffer_minutes": buffer_minutes,
        "findings": findings,
        "slides": slide_reports,
        "method": (
            "speaker-note completeness plus deterministic text-rate diagnostic; "
            "not a human rehearsal"
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate talk timing and speaker notes.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        report = validate(args.pptx, args.plan)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"Timing validation failed: {exc}", file=sys.stderr)
        return 2
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
