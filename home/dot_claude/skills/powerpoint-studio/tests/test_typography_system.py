from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from scripts.audit_typography_system import audit


def make_deck(
    path: Path,
    *,
    body_font: str = "Noto Sans CJK JP",
    title_size: float = 30,
    body_size: float = 16,
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(10), Inches(0.8))
    title.name = "slide-title"
    title_run = title.text_frame.paragraphs[0].add_run()
    title_run.text = "Title"
    title_run.font.name = body_font
    title_run.font.size = Pt(title_size)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(8), Inches(1))
    body.name = "body"
    body_run = body.text_frame.paragraphs[0].add_run()
    body_run.text = "Body"
    body_run.font.name = body_font
    body_run.font.size = Pt(body_size)
    page = slide.shapes.add_textbox(Inches(11.8), Inches(7.12), Inches(0.8), Inches(0.2))
    page.name = "page-number"
    page_run = page.text_frame.paragraphs[0].add_run()
    page_run.text = "01 / 01"
    page_run.font.name = body_font
    page_run.font.size = Pt(11)
    presentation.save(path)


def policy() -> dict:
    return {
        "meta": {
            "typography_policy": {
                "enabled": True,
                "primary_font": "Noto Sans CJK JP",
                "mono_font": "Noto Sans Mono CJK JP",
                "minimum_font_size": 11,
                "deck_title_range": [40, 56],
                "slide_title_range": [25, 34],
                "body_range": [14, 22],
                "caption_range": [11, 13.5],
                "page_number_range": [10, 12],
                "max_size_steps_per_slide": 6,
                "title_minimum_difference_pt": 5,
                "title_minimum_ratio": 1.2,
            }
        }
    }


def test_consistent_typography_passes(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    assert audit(deck, policy())["passed"]


def test_unapproved_font_fails(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck, body_font="Arial")
    report = audit(deck, policy())
    assert not report["passed"]
    assert any(item["category"] == "font-family" for item in report["findings"])


def test_small_body_font_fails(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck, body_size=9)
    report = audit(deck, policy())
    assert not report["passed"]
    assert any(item["category"] == "minimum-font-size" for item in report["findings"])


def test_title_range_is_enforced(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck, title_size=20)
    report = audit(deck, policy())
    assert not report["passed"]
    assert any(item["category"] == "slide-title-size" for item in report["findings"])


def test_policy_is_optional(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    report = audit(deck, {"meta": {}})
    assert report["enabled"] is False
    assert report["passed"]


def test_generated_diagram_font_is_checked(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    diagrams = tmp_path / "assets" / "diagrams"
    diagrams.mkdir(parents=True)
    (diagrams / "flow.dot").write_text(
        'digraph G { node [fontname="Arial", fontsize=16]; a -> b; }\n',
        encoding="utf-8",
    )
    plan = policy()
    plan["meta"]["typography_policy"]["generated_diagram_sources"] = [
        {"path": "assets/diagrams/flow.dot"}
    ]
    report = audit(deck, plan)
    assert not report["passed"]
    assert any(item["category"] == "generated-diagram-font" for item in report["findings"])
