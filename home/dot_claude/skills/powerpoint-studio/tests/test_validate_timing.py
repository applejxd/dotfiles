from pathlib import Path

from pptx import Presentation
from scripts.validate_timing import validate


def test_timing_validation_passes_with_notes(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = (
        "最初に結論を述べる。図の左から右を順番に指す。具体例を一つ説明し、"
        "誤解しやすい点を補足してから、次のページへ接続する。"
    )
    pptx = tmp_path / "talk.pptx"
    presentation.save(pptx)
    plan = tmp_path / "plan.json"
    plan.write_text(
        """
        {
          "meta": {"duration_minutes": 1},
          "slides": [
            {"timing_seconds": 60, "optional_cut_seconds": 10}
          ]
        }
        """,
        encoding="utf-8",
    )
    report = validate(pptx, plan)
    assert report["passed"]


def test_missing_notes_fail(tmp_path: Path) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    pptx = tmp_path / "talk.pptx"
    presentation.save(pptx)
    plan = tmp_path / "plan.json"
    plan.write_text(
        """
        {
          "meta": {"duration_minutes": 1},
          "slides": [
            {"timing_seconds": 60, "optional_cut_seconds": 0}
          ]
        }
        """,
        encoding="utf-8",
    )
    report = validate(pptx, plan)
    assert not report["passed"]
