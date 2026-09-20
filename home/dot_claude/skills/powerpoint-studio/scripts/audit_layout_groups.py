#!/usr/bin/env python3
"""Opt-in, exact-name alignment and effective font-size checks.

meta.layout_checks declares every tolerance and minimum; absent/empty checks
disable the audit. Repeated shape names select every instance, but repeated
selectors never double-count a shape. Font sizes resolve run then paragraph,
not theme/master guesses. Grouped targets and rotated alignment targets are
reported as unsupported rather than comparing misleading local coordinates.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.exc import PackageNotFoundError

AXES = ("left", "right", "top", "bottom", "center-x", "center-y")
TEXT_ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
EMU_PER_INCH = 914400


def _report() -> dict[str, Any]:
    return {
        "enabled": True,
        "passed": True,
        "errors": 0,
        "warnings": 0,
        "findings": [],
        "checks": [],
    }


def _error(report: dict, category: str, message: str, **context: Any) -> None:
    report["findings"].append(
        {"severity": "error", "category": category, "message": message, **context}
    )
    report["errors"] += 1
    report["passed"] = False


def _number(value: Any, *, positive: bool = False) -> bool:
    try:
        return (
            isinstance(value, (int, float))
            and not isinstance(value, bool)
            and math.isfinite(value)
            and (value > 0 if positive else value >= 0)
        )
    except OverflowError:
        return False


def _validate(checks: list, report: dict) -> None:
    for index, item in enumerate(checks, start=1):
        if not isinstance(item, dict):
            _error(report, "input", "Each layout check must be an object.", check_index=index)
            continue
        slide, names, kind = item.get("slide"), item.get("shapes"), item.get("kind")
        problems = []
        if isinstance(slide, bool) or not isinstance(slide, int) or slide < 1:
            problems.append("slide must be a positive 1-based integer")
        if (
            not isinstance(names, list)
            or not names
            or not all(isinstance(name, str) and name.strip() for name in names)
        ):
            problems.append("shapes must be a nonempty list of exact, nonempty names")
        if kind == "alignment":
            if item.get("axis") not in AXES:
                problems.append(f"axis must be one of {AXES}")
            if not _number(item.get("tolerance_inches")):
                problems.append("tolerance_inches must be finite and nonnegative")
            if "text_anchor" in item and (
                not isinstance(item["text_anchor"], str) or item["text_anchor"] not in TEXT_ANCHORS
            ):
                problems.append("text_anchor must be top, middle, or bottom")
        elif kind == "font-size":
            if not _number(item.get("minimum_pt"), positive=True):
                problems.append("minimum_pt must be finite and positive")
            if not isinstance(item.get("require_no_autofit", False), bool):
                problems.append("require_no_autofit must be boolean")
        else:
            problems.append("kind must be alignment or font-size")
        for problem in problems:
            _error(report, "input", problem, check_index=index)


def _walk(shapes: Any, grouped: bool = False):
    for shape in shapes:
        yield shape, grouped
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes, True)


def _alignment(report: dict, item: dict, selected: list, record: dict) -> None:
    if len(selected) < 2:
        _error(
            report,
            "selector",
            "Alignment requires at least two distinct matched shapes.",
            check_index=record["check_index"],
            slide=item["slide"],
        )
        return
    anchors = []
    for shape in selected:
        context = {
            "check_index": record["check_index"],
            "slide": item["slide"],
            "shape": shape.name,
            "shape_id": shape.shape_id,
        }
        if shape.rotation:
            _error(
                report,
                "unsupported-transform",
                "Rotated alignment targets are not supported.",
                **context,
            )
            continue
        if "text_anchor" in item:
            if not shape.has_text_frame:
                _error(report, "shape-type", "text_anchor requires a text frame.", **context)
                continue
            actual = shape.text_frame.vertical_anchor
            if actual != TEXT_ANCHORS[item["text_anchor"]]:
                _error(
                    report,
                    "text-anchor",
                    f"Expected explicit {item['text_anchor']} text anchor; found {actual}.",
                    **context,
                )
        box = (shape.left, shape.top, shape.width, shape.height)
        if not all(isinstance(value, int) for value in box):
            _error(report, "shape-type", "Shape has unresolved geometry.", **context)
            continue
        x, y, width, height = box
        anchor = {
            "left": x,
            "right": x + width,
            "top": y,
            "bottom": y + height,
            "center-x": x + width / 2,
            "center-y": y + height / 2,
        }[item["axis"]] / EMU_PER_INCH
        anchors.append({"shape": shape.name, "shape_id": shape.shape_id, "anchor_inches": anchor})
    record["anchors"] = anchors
    if len(anchors) >= 2:
        values = [anchor["anchor_inches"] for anchor in anchors]
        spread = max(values) - min(values)
        record["spread_inches"] = spread
        if spread > item["tolerance_inches"] + 1e-9:
            _error(
                report,
                "alignment",
                f"{item['axis']} anchors span {spread:.6g}in; "
                f"tolerance is {item['tolerance_inches']:g}in.",
                check_index=record["check_index"],
                slide=item["slide"],
            )


def _font_sizes(report: dict, item: dict, shape: Any, check_index: int) -> list[dict]:
    context = {
        "check_index": check_index,
        "slide": item["slide"],
        "shape": shape.name,
        "shape_id": shape.shape_id,
    }
    if not shape.has_text_frame:
        _error(report, "shape-type", "Font-size checks require text shapes.", **context)
        return []
    norms = shape.text_frame._txBody.bodyPr.xpath("./a:normAutofit")
    if norms and item.get("require_no_autofit", False):
        _error(
            report,
            "autofit",
            "normAutofit is forbidden; PowerPoint may shrink this text later.",
            **context,
        )
    scale = 1.0
    if norms:
        try:
            if len(norms) != 1:
                raise ValueError("multiple normAutofit elements")
            raw = norms[0].get("fontScale", "100000")
            scale = float(raw[:-1]) / 100 if raw.endswith("%") else int(raw) / 100000
            if not _number(scale, positive=True) or scale > 1:
                raise ValueError("fontScale must be greater than 0 and at most 100%")
        except (ValueError, OverflowError) as exc:
            _error(report, "input", f"Invalid normAutofit fontScale: {exc}", **context)
            return []
    measurements = []
    text_seen = False
    for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs, start=1):
        if paragraph._p.xpath("./a:fld"):
            text_seen = True
            _error(
                report,
                "font-size-unresolved",
                "Field text is not supported by the run/paragraph font resolver.",
                paragraph=paragraph_index,
                **context,
            )
        for run_index, run in enumerate(paragraph.runs, start=1):
            if not run.text.strip():
                continue
            text_seen = True
            size = run.font.size
            if size is None:
                size = paragraph.font.size
            if size is None:
                _error(
                    report,
                    "font-size-unresolved",
                    "No explicit run or paragraph font size.",
                    paragraph=paragraph_index,
                    run=run_index,
                    **context,
                )
                continue
            stored = size.pt
            if not _number(stored, positive=True):
                _error(report, "input", "Stored font size must be finite and positive.", **context)
                continue
            effective = stored * scale
            measurements.append(
                {
                    **context,
                    "paragraph": paragraph_index,
                    "run": run_index,
                    "stored_pt": stored,
                    "font_scale": scale,
                    "effective_pt": effective,
                }
            )
            if effective < item["minimum_pt"] - 1e-9:
                _error(
                    report,
                    "font-size",
                    f"Effective font is {effective:g}pt ({stored:g}pt x {scale:g}); "
                    f"minimum is {item['minimum_pt']:g}pt.",
                    paragraph=paragraph_index,
                    run=run_index,
                    **context,
                )
    if not text_seen:
        _error(
            report,
            "font-size-unresolved",
            "Selected text shape contains no measurable text.",
            **context,
        )
    return measurements


def audit(pptx_path: Path, plan: Any) -> dict[str, Any]:
    report = _report()
    if not isinstance(plan, dict) or not isinstance(plan.get("meta", {}), dict):
        _error(report, "input", "Plan and meta must be JSON objects.")
        return report
    checks = plan.get("meta", {}).get("layout_checks", [])
    if not isinstance(checks, list):
        _error(report, "input", "meta.layout_checks must be an array.")
        return report
    if not checks:
        report["enabled"] = False
        return report
    _validate(checks, report)
    if not report["passed"]:
        return report
    try:
        presentation = Presentation(pptx_path)
        for index, item in enumerate(checks, start=1):
            context = {"check_index": index, "slide": item["slide"]}
            record = {**context, "kind": item["kind"], "matched_shapes": 0}
            report["checks"].append(record)
            if item["slide"] > len(presentation.slides):
                _error(report, "selector", f"Slide {item['slide']} does not exist.", **context)
                continue
            names = set(item["shapes"])
            matches = [
                (shape, grouped)
                for shape, grouped in _walk(presentation.slides[item["slide"] - 1].shapes)
                if shape.name in names
            ]
            record["matched_shapes"] = len(matches)
            for missing in sorted(names - {shape.name for shape, _ in matches}):
                _error(report, "selector", f"No shape matches exact name {missing!r}.", **context)
            selected = []
            for shape, grouped in matches:
                if grouped:
                    _error(
                        report,
                        "unsupported-transform",
                        "Grouped targets require slide-space transforms and are not supported.",
                        shape=shape.name,
                        shape_id=shape.shape_id,
                        **context,
                    )
                else:
                    selected.append(shape)
            if item["kind"] == "alignment":
                _alignment(report, item, selected, record)
            else:
                fonts = [
                    measurement
                    for shape in selected
                    for measurement in _font_sizes(report, item, shape, index)
                ]
                record["fonts"] = fonts
                record["minimum_effective_pt"] = min(
                    (font["effective_pt"] for font in fonts), default=None
                )
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError, XMLSyntaxError) as exc:
        _error(report, "input", f"Cannot inspect PPTX {pptx_path}: {exc}")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.output and args.output.resolve() in (args.pptx.resolve(), args.plan.resolve()):
        parser.exit(1, "Report must not overwrite the PPTX or plan.\n")
    try:
        report = audit(args.pptx, json.loads(args.plan.read_bytes()))
    except (OSError, ValueError) as exc:
        report = _report()
        _error(report, "input", str(exc))
    text = json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    if args.output:
        try:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(text, encoding="utf-8")
        except OSError as exc:
            parser.exit(1, f"Cannot write layout-group report: {exc}\n")
    print(text, end="")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
