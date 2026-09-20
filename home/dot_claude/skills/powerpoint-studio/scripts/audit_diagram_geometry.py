#!/usr/bin/env python3
"""Audit a *.layout.json, optionally binding it to an actual PPTX picture.

Relative image/layout paths in a plan are resolved by validate_deck against the
plan directory. The CLI's positional and --image paths are relative to cwd.
Reports hash the exact input bytes, not a JSON reserialization.

Image audits require --provenance: a build-time JSON manifest containing
layout_sha256, svg_sha256, image_sha256, renderer_sha256. The SVG is the image
path with its extension replaced by .svg; renderer_sha256 hashes the sibling
diagram_geometry.py, which contains both geometry and SVG serialization.
Write the manifest only after successful SVG rendering and rasterization.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

if __package__:
    from .diagram_geometry import EPSILON, add_error, audit, empty_report, svg_from_report
else:
    from diagram_geometry import EPSILON, add_error, audit, empty_report, svg_from_report

PLACEMENT_TOLERANCE_INCHES = 0.02


def read_layout(path: Path) -> tuple[Any, bytes]:
    raw = path.read_bytes()
    return json.loads(raw), raw


def _positive(value: Any) -> bool:
    try:
        return (
            not isinstance(value, bool)
            and isinstance(value, (int, float))
            and math.isfinite(value)
            and value > 0
        )
    except OverflowError:
        return False


def _policy(report: dict[str, Any], policy: Any, minimum: float) -> float:
    if policy is None:
        return minimum
    if not isinstance(policy, dict):
        add_error(report, "input", "typography_policy must be an object.")
        return minimum
    if policy.get("enabled") is not True:
        return minimum
    allowed = [policy.get("primary_font"), policy.get("mono_font")]
    if any(value is not None and not isinstance(value, str) for value in allowed):
        add_error(report, "input", "Typography font names must be strings.")
        return minimum
    allowed = [value.strip() for value in allowed if value and value.strip()]
    if not allowed or report["fonts"].get("font_family") not in allowed:
        add_error(
            report,
            "font-family",
            f"Diagram font {report['fonts'].get('font_family')!r} "
            f"is outside the allowed fonts {allowed!r}.",
        )
    policy_minimum = policy.get("minimum_font_size", 11)
    if not _positive(policy_minimum):
        add_error(
            report, "input", "typography_policy.minimum_font_size must be finite and positive."
        )
    else:
        minimum = max(minimum, policy_minimum)
    return minimum


def _find_shapes(shapes: Any, name: str, *, grouped: bool = False) -> list[tuple[Any, bool]]:
    found = []
    for shape in shapes:
        if shape.name == name:
            found.append((shape, grouped))
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found.extend(_find_shapes(shape.shapes, name, grouped=True))
    return found


def _font_range(report: dict[str, Any], policy: Any, sizes: list[float], context: str) -> None:
    if not isinstance(policy, dict) or policy.get("enabled") is not True:
        return
    bounds = policy.get("generated_diagram_apparent_font_range")
    if bounds is None:
        return
    if (
        not isinstance(bounds, list)
        or len(bounds) != 2
        or not all(_positive(x) for x in bounds)
        or bounds[0] > bounds[1]
    ):
        add_error(
            report,
            "input",
            "generated_diagram_apparent_font_range must be [positive minimum, maximum].",
        )
        return
    for size in sizes:
        if size < bounds[0] - EPSILON or size > bounds[1] + EPSILON:
            add_error(
                report,
                "generated-diagram-apparent-font-size",
                f"{context} text appears at {size:g}pt, outside the typography range {bounds}.",
            )


def _provenance(
    report: dict[str, Any],
    image_path: Path | None,
    provenance_path: Path | None,
    expected_svg_sha256: str | None,
) -> None:
    if image_path is None:
        if provenance_path is not None:
            add_error(
                report, "provenance", "Provenance verification requires an explicit image path."
            )
        return
    if provenance_path is None:
        add_error(report, "provenance", "Image audits require a build-time provenance manifest.")
        return
    svg_path = image_path.with_suffix(".svg")
    renderer_path = Path(__file__).resolve().with_name("diagram_geometry.py")
    artifacts = {"provenance": provenance_path, "svg": svg_path, "renderer": renderer_path}
    contents = {}
    for key, path in artifacts.items():
        report[key] = str(path.resolve())
        try:
            contents[key] = path.read_bytes()
            report["sha256"][key] = hashlib.sha256(contents[key]).hexdigest()
        except OSError as exc:
            add_error(report, "provenance", f"Cannot read {key} artifact {path}: {exc}")
    if "provenance" not in contents:
        return
    try:
        manifest = json.loads(contents["provenance"])
    except ValueError as exc:
        add_error(report, "provenance", f"Invalid provenance JSON: {exc}")
        return
    if not isinstance(manifest, dict):
        add_error(report, "provenance", "Provenance root must be a JSON object.")
        return
    for key in ("layout", "svg", "image", "renderer"):
        expected = manifest.get(f"{key}_sha256")
        if not isinstance(expected, str) or re.fullmatch(r"[0-9a-fA-F]{64}", expected) is None:
            add_error(
                report, "provenance", f"Provenance {key}_sha256 must be a 64-digit SHA256 string."
            )
        elif report["sha256"].get(key) != expected.lower():
            add_error(
                report,
                "provenance",
                f"Provenance {key}_sha256 does not match the current {key}; rebuild SVG and image.",
            )
    if expected_svg_sha256 is not None:
        report["expected_svg_sha256"] = expected_svg_sha256
        if report["sha256"].get("svg") != expected_svg_sha256:
            add_error(
                report,
                "provenance",
                "SVG does not match the deterministic output of the current layout and renderer; "
                "rebuild SVG and image.",
            )


def audit_file(
    layout_path: Path,
    *,
    image_path: Path | None = None,
    provenance_path: Path | None = None,
    pptx_path: Path | None = None,
    slide: int | None = None,
    shape_name: str | None = None,
    typography_policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Verify build-time provenance before accepting image and actual PPTX bindings."""
    report = empty_report()
    try:
        data, raw = read_layout(layout_path)
    except (OSError, ValueError) as exc:
        add_error(report, "input", f"Cannot read layout {layout_path}: {exc}")
        return report
    report = audit(data)
    report["layout"] = str(layout_path.resolve())
    report["sha256"] = {"layout": hashlib.sha256(raw).hexdigest()}
    if not report["fonts"]:
        return report
    expected_svg_sha256 = (
        hashlib.sha256(svg_from_report(report).encode("utf-8")).hexdigest()
        if report["passed"] and image_path is not None
        else None
    )
    minimum = _policy(report, typography_policy, float(data["placement"]["min_font_pt"]))
    source_sizes = report["fonts"]["font_sizes"]
    declared_sizes = report["fonts"]["apparent_font_sizes_pt"]
    _font_range(report, typography_policy, declared_sizes, "Declared placement")
    for size in declared_sizes:
        if size < minimum - EPSILON and minimum > data["placement"]["min_font_pt"]:
            add_error(
                report,
                "minimum-font-size",
                f"Declared placement gives {size:g}pt text, "
                f"below typography minimum {minimum:g}pt.",
            )
    image_raw = None
    if image_path is not None:
        report["image"] = str(image_path.resolve())
        try:
            image_raw = image_path.read_bytes()
            report["sha256"]["image"] = hashlib.sha256(image_raw).hexdigest()
            with Image.open(io.BytesIO(image_raw)) as image:
                image.verify()
            with Image.open(io.BytesIO(image_raw)) as image:
                image_width, image_height = image.size
            report["image_dimensions"] = {"width": image_width, "height": image_height}
            expected_ratio = data["width"] / data["height"]
            # A single pixel of raster rounding is allowed, never arbitrary padding/cropping.
            if abs(image_width - image_height * expected_ratio) > max(1, expected_ratio) + EPSILON:
                add_error(
                    report,
                    "image-aspect",
                    "Specified image aspect ratio does not match the layout canvas.",
                )
        except (OSError, ValueError, Image.DecompressionBombError) as exc:
            add_error(report, "image-content", f"Cannot read the specified image: {exc}")
    _provenance(report, image_path, provenance_path, expected_svg_sha256)
    if pptx_path is None:
        if slide is not None or shape_name is not None:
            add_error(report, "input", "--slide and --shape require --pptx.")
        return report
    if image_path is None:
        add_error(report, "input", "PPTX binding requires an explicit image path.")
    if isinstance(slide, bool) or not isinstance(slide, int) or slide < 1:
        add_error(report, "input", "slide must be a positive 1-based integer.")
        return report
    if not isinstance(shape_name, str) or not shape_name.strip():
        add_error(report, "input", "PPTX binding requires a nonempty shape name.")
        return report
    report.update(pptx=str(pptx_path.resolve()), slide=slide, shape=shape_name)
    try:
        pptx_raw = pptx_path.read_bytes()
        report["sha256"]["pptx"] = hashlib.sha256(pptx_raw).hexdigest()
        presentation = Presentation(io.BytesIO(pptx_raw))
        if slide > len(presentation.slides):
            add_error(report, "image-placement", f"PPTX has no slide {slide}.")
            return report
        matches = _find_shapes(presentation.slides[slide - 1].shapes, shape_name)
        if len(matches) != 1:
            add_error(
                report,
                "image-placement",
                f"Expected exactly one named picture {shape_name!r} on slide {slide}; "
                f"found {len(matches)}.",
            )
            return report
        picture, grouped = matches[0]
        if picture.shape_type != MSO_SHAPE_TYPE.PICTURE:
            add_error(report, "image-content", f"Shape {shape_name!r} is not an embedded picture.")
            return report
        report["embedded_image_sha256"] = hashlib.sha256(picture.image.blob).hexdigest()
        if image_raw is not None and picture.image.blob != image_raw:
            add_error(
                report,
                "image-content",
                "PPTX embedded image differs from the specified image (stale or substituted).",
            )
        if (
            grouped
            or picture.rotation != 0
            or any(
                abs(value) > EPSILON
                for value in (
                    picture.crop_left,
                    picture.crop_right,
                    picture.crop_top,
                    picture.crop_bottom,
                )
            )
        ):
            add_error(
                report,
                "image-placement",
                "Diagram pictures must be ungrouped, unrotated, and uncropped.",
            )
        transform = picture._element.spPr.xfrm
        if transform is not None and any(
            transform.get(key) in ("1", "true") for key in ("flipH", "flipV")
        ):
            add_error(report, "image-placement", "Flipped diagram pictures are not supported.")
        actual_width, actual_height = picture.width / 914400, picture.height / 914400
        scale = min(72 * actual_width / data["width"], 72 * actual_height / data["height"])
        actual_sizes = [size * scale for size in source_sizes]
        _font_range(report, typography_policy, actual_sizes, "Actual PPTX placement")
        report["actual_placement"] = {
            "width_inches": actual_width,
            "height_inches": actual_height,
            "scale_pt_per_unit": scale,
            "apparent_font_sizes_pt": actual_sizes,
            "minimum_apparent_font_pt": min(actual_sizes) if actual_sizes else None,
            "tolerance_inches": PLACEMENT_TOLERANCE_INCHES,
        }
        dimensions = report["dimensions"]
        if any(
            abs(actual - expected) > PLACEMENT_TOLERANCE_INCHES
            for actual, expected in (
                (actual_width, dimensions["expected_width_inches"]),
                (actual_height, dimensions["expected_height_inches"]),
            )
        ):
            add_error(
                report,
                "image-placement",
                f"Actual picture is {actual_width:g} x {actual_height:g}in, expected contain size "
                f"{dimensions['expected_width_inches']:g} x "
                f"{dimensions['expected_height_inches']:g}in.",
            )
        if (
            picture.left < 0
            or picture.top < 0
            or (
                picture.left + picture.width > presentation.slide_width
                or picture.top + picture.height > presentation.slide_height
            )
        ):
            add_error(report, "image-placement", "Diagram picture extends beyond the slide canvas.")
        for size in actual_sizes:
            if size < minimum - EPSILON:
                add_error(
                    report,
                    "minimum-font-size",
                    f"Actual PPTX placement gives {size:g}pt text, below minimum {minimum:g}pt.",
                )
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError, XMLSyntaxError) as exc:
        add_error(report, "input", f"Cannot inspect PPTX {pptx_path}: {exc}")
    return report


def write_report(report: dict[str, Any], path: Path | None) -> None:
    text = json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    if path is not None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
    print(text, end="")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("layout", type=Path)
    parser.add_argument("--output", "--report", type=Path)
    parser.add_argument("--image", type=Path)
    parser.add_argument(
        "--provenance", type=Path, help="Required build-time hash manifest for image audits."
    )
    parser.add_argument("--pptx", type=Path)
    parser.add_argument("--slide", type=int)
    parser.add_argument("--shape")
    parser.add_argument("--plan", type=Path, help="Read meta.typography_policy from this plan.")
    args = parser.parse_args()
    if args.output and args.output.resolve() in {
        path.resolve()
        for path in (
            args.layout,
            args.image,
            args.pptx,
            args.plan,
            args.provenance,
            args.image.with_suffix(".svg") if args.image else None,
            Path(__file__).resolve().with_name("diagram_geometry.py"),
        )
        if path is not None
    }:
        parser.exit(1, "Report output must not overwrite an input artifact.\n")
    try:
        policy = None
        if args.plan:
            plan = json.loads(args.plan.read_bytes())
            if not isinstance(plan, dict) or not isinstance(plan.get("meta", {}), dict):
                raise ValueError("Plan and meta must be objects.")
            policy = plan.get("meta", {}).get("typography_policy")
        report = audit_file(
            args.layout,
            image_path=args.image,
            pptx_path=args.pptx,
            slide=args.slide,
            shape_name=args.shape,
            typography_policy=policy,
            provenance_path=args.provenance,
        )
    except (OSError, ValueError) as exc:
        report = empty_report()
        add_error(report, "input", str(exc))
    try:
        write_report(report, args.output)
    except OSError as exc:
        parser.exit(1, f"Cannot write geometry report: {exc}\n")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
