import hashlib
import json
import math
import subprocess
import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from scripts.audit_layout_groups import audit

from scripts import validate_deck

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


@pytest.fixture
def deck(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, name in enumerate(("header", "cell", "cell")):
        shape = slide.shapes.add_textbox(Inches(5.13), Inches(index + 1), Inches(2), Inches(0.5))
        shape.name = name
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = f"Label {index}"
        run.font.size = Pt(14)
    path = tmp_path / "deck.pptx"
    presentation.save(path)
    return presentation, path


def alignment(**changes):
    return {
        "slide": 1,
        "kind": "alignment",
        "shapes": ["header", "cell"],
        "axis": "left",
        "tolerance_inches": 0.03,
        **changes,
    }


def font_check(**changes):
    return {"slide": 1, "kind": "font-size", "shapes": ["cell"], "minimum_pt": 14, **changes}


def check(deck, *checks):
    presentation, path = deck
    presentation.save(path)
    report = audit(path, {"meta": {"layout_checks": list(checks)}})
    assert report["errors"] == len(report["findings"])
    assert report["passed"] == (report["errors"] == 0)
    json.dumps(report, allow_nan=False)
    return report


def categories(report):
    return {finding["category"] for finding in report["findings"]}


@pytest.mark.parametrize("plan", [{}, {"meta": {}}, {"meta": {"layout_checks": []}}])
def test_absent_checks_disabled_without_loading_pptx(tmp_path, plan):
    report = audit(tmp_path / "absent.pptx", plan)
    assert report["enabled"] is False
    assert report["passed"] and report["errors"] == 0


def test_left_alignment_checks_all_repeated_names_and_exact_regression(deck):
    assert check(deck, alignment())["passed"]
    presentation, _ = deck
    presentation.slides[0].shapes[0].left = Inches(5.9)
    report = check(deck, alignment())
    assert "alignment" in categories(report)
    assert report["checks"][0]["spread_inches"] == pytest.approx(0.77)
    assert report["checks"][0]["matched_shapes"] == 3
    presentation.slides[0].shapes[0].left = Inches(5.13)
    presentation.slides[0].shapes[2].left = Inches(5.17)
    assert "alignment" in categories(check(deck, alignment()))


@pytest.mark.parametrize("axis", ["left", "right", "top", "bottom", "center-x", "center-y"])
def test_all_alignment_anchors_and_tolerance(deck, axis):
    presentation, _ = deck
    for shape in presentation.slides[0].shapes:
        shape.top = Inches(1)
    assert check(deck, alignment(axis=axis))["passed"]
    moved = presentation.slides[0].shapes[2]
    attribute = "left" if axis in ("left", "right", "center-x") else "top"
    setattr(moved, attribute, getattr(moved, attribute) + Inches(0.03))
    assert check(deck, alignment(axis=axis))["passed"]
    setattr(moved, attribute, getattr(moved, attribute) + Inches(0.001))
    assert "alignment" in categories(check(deck, alignment(axis=axis)))


def test_text_anchor_rejects_middle_aligned_lists_at_the_same_top(deck):
    presentation, _ = deck
    for shape in presentation.slides[0].shapes:
        shape.top = Inches(1)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    assert check(deck, alignment(axis="top", text_anchor="top"))["passed"]
    last = presentation.slides[0].shapes[2]
    last.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    report = check(deck, alignment(axis="top", text_anchor="top"))
    assert "text-anchor" in categories(report)
    last.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    assert check(deck, alignment(axis="top", text_anchor="top"))["passed"]


@pytest.mark.parametrize(
    "name,value",
    [
        ("top", MSO_ANCHOR.TOP),
        ("middle", MSO_ANCHOR.MIDDLE),
        ("bottom", MSO_ANCHOR.BOTTOM),
    ],
)
def test_text_anchor_accepts_declared_vertical_alignment(deck, name, value):
    presentation, _ = deck
    for shape in presentation.slides[0].shapes:
        shape.text_frame.vertical_anchor = value
    assert check(deck, alignment(text_anchor=name))["passed"]


@pytest.mark.parametrize("value", ["side", None, 42, {}, [], True])
def test_text_anchor_rejects_invalid_declarations(deck, value):
    assert "input" in categories(check(deck, alignment(text_anchor=value)))


def test_text_anchor_requires_a_text_frame(deck):
    presentation, path = deck
    image = path.with_suffix(".png")
    Image.new("RGB", (10, 10), "white").save(image)
    picture = presentation.slides[0].shapes.add_picture(
        str(image), Inches(5.13), Inches(1), Inches(1), Inches(1)
    )
    picture.name = "picture"
    report = check(deck, alignment(shapes=["header", "picture"], text_anchor="top"))
    assert "shape-type" in categories(report)


def test_text_anchor_does_not_guess_inherited_alignment(deck):
    assert "text-anchor" in categories(check(deck, alignment(text_anchor="top")))


@pytest.mark.parametrize(
    "axis,position,extent,fraction",
    [
        ("right", "left", "width", 1),
        ("center-x", "left", "width", 0.5),
        ("bottom", "top", "height", 1),
        ("center-y", "top", "height", 0.5),
    ],
)
def test_anchors_use_extents_not_just_origins(deck, axis, position, extent, fraction):
    presentation, _ = deck
    for shape in presentation.slides[0].shapes:
        shape.top = Inches(1)
    shape = presentation.slides[0].shapes[2]
    setattr(shape, extent, getattr(shape, extent) + Inches(0.2))
    setattr(shape, position, getattr(shape, position) - Inches(0.2 * fraction))
    assert check(deck, alignment(axis=axis))["passed"]
    setattr(shape, extent, getattr(shape, extent) + Inches(0.2))
    assert "alignment" in categories(check(deck, alignment(axis=axis)))


def test_selector_duplicates_not_double_counted(deck):
    assert "selector" in categories(check(deck, alignment(shapes=["header", "header"])))
    report = check(deck, alignment(shapes=["cell", "cell"]))
    assert report["passed"]
    assert report["checks"][0]["matched_shapes"] == 2


def test_every_selector_must_match_exactly(deck):
    assert "selector" in categories(check(deck, alignment(shapes=["cell", "missing"])))
    assert "selector" in categories(check(deck, font_check(shapes=["Cell"])))


def test_font_size_checks_every_repeated_cell_and_run(deck):
    assert check(deck, font_check())["passed"]
    presentation, _ = deck
    paragraph = presentation.slides[0].shapes[2].text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = " smaller"
    run.font.size = Pt(11.5)
    report = check(deck, font_check())
    assert "font-size" in categories(report)
    assert report["checks"][0]["minimum_effective_pt"] == pytest.approx(11.5)
    assert report["checks"][0]["matched_shapes"] == 2
    assert check(deck, font_check(minimum_pt=11.5))["passed"]


def test_paragraph_fallback_and_explicit_run_precedence(deck):
    presentation, _ = deck
    for shape in presentation.slides[0].shapes:
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.font.size = Pt(15)
        paragraph.runs[0].font.size = None
    assert check(deck, font_check())["passed"]
    presentation.slides[0].shapes[2].text_frame.paragraphs[0].runs[0].font.size = Pt(11.5)
    assert "font-size" in categories(check(deck, font_check()))


def test_unresolved_font_and_empty_text_are_not_ignored(deck):
    presentation, _ = deck
    paragraph = presentation.slides[0].shapes[2].text_frame.paragraphs[0]
    paragraph.runs[0].font.size = None
    assert "font-size-unresolved" in categories(check(deck, font_check()))
    paragraph.clear()
    assert "font-size-unresolved" in categories(check(deck, font_check()))


def norm_autofit(shape, scale=None):
    body = shape.text_frame._txBody.bodyPr
    for child in list(body):
        if child.tag.endswith(("normAutofit", "spAutoFit", "noAutofit")):
            body.remove(child)
    normal = OxmlElement("a:normAutofit")
    if scale is not None:
        normal.set("fontScale", scale)
    body.append(normal)


@pytest.mark.parametrize("scale", ["80000", "80%"])
def test_explicit_autofit_scale_reduces_effective_size(deck, scale):
    presentation, _ = deck
    shape = presentation.slides[0].shapes[2]
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    norm_autofit(shape, scale)
    report = check(deck, font_check())
    assert "font-size" in categories(report)
    assert report["checks"][0]["minimum_effective_pt"] == pytest.approx(12.8)


@pytest.mark.parametrize("scale", [None, "100000", "100%"])
def test_require_no_autofit_rejects_even_currently_unshrunk_text(deck, scale):
    presentation, _ = deck
    norm_autofit(presentation.slides[0].shapes[2], scale)
    assert check(deck, font_check())["passed"]
    assert "autofit" in categories(check(deck, font_check(require_no_autofit=True)))


@pytest.mark.parametrize("scale", ["NaN", "-1", "0", "100001", "oops", "Infinity", "101%"])
def test_invalid_autofit_scale_fails(deck, scale):
    presentation, _ = deck
    norm_autofit(presentation.slides[0].shapes[2], scale)
    assert not check(deck, font_check())["passed"]


def test_selected_picture_is_not_a_text_shape(deck, tmp_path):
    presentation, _ = deck
    image = tmp_path / "image.png"
    Image.new("RGB", (20, 20)).save(image)
    picture = presentation.slides[0].shapes.add_picture(str(image), 0, 0)
    picture.name = "cell"
    assert "shape-type" in categories(check(deck, font_check()))


@pytest.mark.parametrize(
    "declaration",
    [
        None,
        {},
        "bad",
        alignment(kind="spacing"),
        alignment(axis="baseline"),
        alignment(axis=[]),
        alignment(slide=None),
        alignment(slide=0),
        alignment(slide=True),
        alignment(slide=2.1),
        alignment(shapes=[]),
        alignment(shapes="cell"),
        alignment(shapes=[1]),
        alignment(tolerance_inches=math.nan),
        alignment(tolerance_inches=math.inf),
        alignment(tolerance_inches=-0.01),
        alignment(tolerance_inches=True),
        font_check(minimum_pt=0),
        font_check(minimum_pt=-1),
        font_check(minimum_pt=math.nan),
        font_check(minimum_pt="14"),
        font_check(require_no_autofit="true"),
    ],
)
def test_invalid_declarations_fail_without_crashing(deck, declaration):
    report = check(deck, declaration)
    assert report["enabled"] and "input" in categories(report)


@pytest.mark.parametrize("checks", [None, {}, "", False])
def test_invalid_collection_is_not_disabled(deck, checks):
    _, path = deck
    report = audit(path, {"meta": {"layout_checks": checks}})
    assert report["enabled"] and not report["passed"]


def test_out_of_range_slide_and_missing_fields_fail(deck):
    assert "selector" in categories(check(deck, alignment(slide=2)))
    for field in ("slide", "kind", "shapes", "axis", "tolerance_inches"):
        declaration = alignment()
        del declaration[field]
        assert "input" in categories(check(deck, declaration))


@pytest.mark.parametrize("plan", [None, [], {"meta": None}, {"meta": []}])
def test_invalid_plan_root_fails(deck, plan):
    _, path = deck
    assert "input" in categories(audit(path, plan))


def test_grouped_targets_are_reported_instead_of_silently_skipped(deck):
    presentation, _ = deck
    group = presentation.slides[0].shapes.add_group_shape()
    shape = group.shapes.add_textbox(0, 0, Inches(1), Inches(1))
    shape.name = "cell"
    shape.text = "Grouped"
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    report = check(deck, alignment(), font_check())
    assert "unsupported-transform" in categories(report)
    assert all(item["matched_shapes"] == 4 for item in report["checks"][:1])
    assert report["checks"][1]["matched_shapes"] == 3


def test_cli_writes_pass_and_fail_reports(deck, tmp_path):
    _, path = deck
    plan = tmp_path / "plan.json"
    output = tmp_path / "report.json"
    for declarations, expected in [([alignment()], 0), ([alignment(axis="top")], 1), (None, 1)]:
        plan.write_text(json.dumps({"meta": {"layout_checks": declarations}}))
        result = subprocess.run(
            [
                sys.executable,
                str(SCRIPTS / "audit_layout_groups.py"),
                str(path),
                "--plan",
                str(plan),
                "--output",
                str(output),
            ],
            capture_output=True,
        )
        assert result.returncode == expected, result.stderr
        assert json.loads(output.read_bytes())["passed"] == (expected == 0)


@pytest.mark.parametrize("declared", [True, False])
def test_validate_deck_integration_and_report_freshness(deck, tmp_path, monkeypatch, declared):
    _, path = deck
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(json.dumps({"meta": {"layout_checks": [alignment()]}} if declared else {}))
    output = tmp_path / "qa"
    original_run = validate_deck.run
    calls = []

    def run(command, env, stdout_path=None):
        script = Path(command[1]).name
        calls.append(script)
        if script == "audit_layout_groups.py":
            original_run(command, env, stdout_path)
        elif script == "quality_gate.py":
            Path(command[command.index("--output") + 1]).write_text(json.dumps({"slide_count": 1}))
        elif script == "render_slides.py":
            Image.new("RGB", (20, 20)).save(output / "rendered" / "slide-1.png")
        elif script == "detect_font.py":
            stdout_path.write_text("{}")

    monkeypatch.setattr(validate_deck, "run", run)
    monkeypatch.setattr(
        sys,
        "argv",
        ["validate_deck.py", str(path), "--plan", str(plan_path), "--output-dir", str(output)],
    )
    assert validate_deck.main() == 0
    assert calls.count("audit_layout_groups.py") == 1
    summary_path = output / "qa-summary.json"
    summary = json.loads(summary_path.read_bytes())
    report_path = Path(summary["layout_groups_report"])
    assert json.loads(report_path.read_bytes())["enabled"] is declared
    assert (
        summary["sha256"]["layout_groups_report"]
        == hashlib.sha256(report_path.read_bytes()).hexdigest()
    )
    command = [sys.executable, str(SCRIPTS / "verify_qa_freshness.py"), str(summary_path)]
    assert subprocess.run(command, capture_output=True).returncode == 0
    report_path.write_bytes(report_path.read_bytes() + b"\n")
    assert subprocess.run(command, capture_output=True).returncode != 0
    report_path.unlink()
    assert subprocess.run(command, capture_output=True).returncode != 0
    summary["sha256"].pop("layout_groups_report")
    summary_path.write_text(json.dumps(summary))
    assert subprocess.run(command, capture_output=True).returncode != 0


def test_failed_layout_check_stops_deck_gate_and_preserves_failed_report(
    deck, tmp_path, monkeypatch
):
    _, path = deck
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(json.dumps({"meta": {"layout_checks": [alignment(axis="top")]}}))
    output = tmp_path / "qa"
    original_run = validate_deck.run
    calls = []

    def run(command, env, stdout_path=None):
        calls.append(Path(command[1]).name)
        assert calls == ["audit_layout_groups.py"]
        original_run(command, env, stdout_path)

    monkeypatch.setattr(validate_deck, "run", run)
    monkeypatch.setattr(
        sys,
        "argv",
        ["validate_deck.py", str(path), "--plan", str(plan_path), "--output-dir", str(output)],
    )
    assert validate_deck.main() == 1
    assert not json.loads((output / "layout-groups-report.json").read_bytes())["passed"]
    assert not (output / "qa-summary.json").exists()


@pytest.mark.parametrize(
    "summary,expected",
    [
        ({}, 0),
        ({"layout_groups_report": None, "sha256": {"layout_groups_report": None}}, 0),
        ({"sha256": {"layout_groups_report": "missing-path"}}, 1),
    ],
)
def test_freshness_optional_layout_report_compatibility(tmp_path, summary, expected):
    path = tmp_path / "summary.json"
    path.write_text(json.dumps(summary))
    result = subprocess.run(
        [sys.executable, str(SCRIPTS / "verify_qa_freshness.py"), str(path)], capture_output=True
    )
    assert result.returncode == expected
