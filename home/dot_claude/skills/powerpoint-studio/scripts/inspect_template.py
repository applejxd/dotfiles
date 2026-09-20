#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

EMU_PER_INCH = 914_400
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def inches(value: int) -> float:
    return round(value / EMU_PER_INCH, 4)


def shape_record(shape: Any) -> dict[str, Any]:
    record: dict[str, Any] = {
        "name": shape.name,
        "shape_type": str(shape.shape_type),
        "bounds": {
            "x": inches(shape.left),
            "y": inches(shape.top),
            "w": inches(shape.width),
            "h": inches(shape.height),
        },
        "is_placeholder": bool(shape.is_placeholder),
    }
    if shape.is_placeholder:
        fmt = shape.placeholder_format
        record["placeholder"] = {
            "idx": fmt.idx,
            "type": str(fmt.type),
        }
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        record["text"] = text
        sizes: list[float] = []
        fonts: set[str] = set()
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size:
                    sizes.append(round(run.font.size.pt, 2))
                if run.font.name:
                    fonts.add(run.font.name)
        record["font_sizes_pt"] = sorted(set(sizes))
        record["fonts"] = sorted(fonts)
    if getattr(shape, "shape_type", None) == 13:
        try:
            width_px, height_px = shape.image.size
            record["image"] = {
                "px": [width_px, height_px],
                "aspect_ratio": round(width_px / height_px, 4),
            }
        except (AttributeError, ValueError, OSError):
            record["image"] = {"error": "unreadable image relation"}
    return record


def theme_fonts(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {"major": {}, "minor": {}}
    with ZipFile(path) as archive:
        names = archive.namelist()
        theme_name = next(
            (name for name in names if re.fullmatch(r"ppt/theme/theme\d+\.xml", name)),
            None,
        )
        if not theme_name:
            return result
        root = etree.fromstring(archive.read(theme_name))
        for role, xpath in (
            ("major", ".//a:themeElements/a:fontScheme/a:majorFont"),
            ("minor", ".//a:themeElements/a:fontScheme/a:minorFont"),
        ):
            node = root.find(xpath, namespaces=NS)
            if node is None:
                continue
            for child in node:
                name = etree.QName(child).localname
                typeface = child.get("typeface")
                script = child.get("script")
                if typeface:
                    result[role][script or name] = typeface
    return result


def inspect(path: Path) -> dict[str, Any]:
    presentation = Presentation(path)
    layouts: list[dict[str, Any]] = []
    for index, layout in enumerate(presentation.slide_layouts):
        placeholders = []
        for shape in layout.placeholders:
            placeholders.append(shape_record(shape))
        layouts.append(
            {
                "index": index,
                "name": layout.name,
                "master_name": layout.slide_master.name,
                "placeholders": placeholders,
            }
        )

    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        slides.append(
            {
                "number": index,
                "layout_name": slide.slide_layout.name,
                "shapes": [shape_record(shape) for shape in slide.shapes],
            }
        )
    return {
        "path": str(path.resolve()),
        "slide_size_inches": {
            "width": inches(presentation.slide_width),
            "height": inches(presentation.slide_height),
        },
        "theme_fonts": theme_fonts(path),
        "masters": [master.name for master in presentation.slide_masters],
        "layouts": layouts,
        "slides": slides,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Inventory a PowerPoint template.")
    parser.add_argument("template", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    data = inspect(args.template)
    rendered = json.dumps(data, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
        print(f"Wrote template inventory: {args.output}")
    else:
        print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
