#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation


def load_json(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("The plan root must be a JSON object.")
    return data


def in_range(value: float, bounds: list[float]) -> bool:
    return float(bounds[0]) <= value <= float(bounds[1])


def audit(pptx_path: Path, plan: dict[str, Any]) -> dict[str, Any]:
    meta = plan.get("meta")
    policy = meta.get("typography_policy", {}) if isinstance(meta, dict) else {}
    enabled = isinstance(policy, dict) and policy.get("enabled") is True
    if not enabled:
        return {
            "enabled": False,
            "passed": True,
            "errors": 0,
            "warnings": 0,
            "findings": [],
        }

    primary = str(policy.get("primary_font", "")).strip()
    mono = str(policy.get("mono_font", "")).strip()
    allowed = {font for font in (primary, mono) if font}
    minimum = float(policy.get("minimum_font_size", 11))
    title_range = list(policy.get("slide_title_range", [25, 34]))
    deck_title_range = list(policy.get("deck_title_range", [40, 56]))
    page_range = list(policy.get("page_number_range", [10, 12]))
    max_steps = int(policy.get("max_size_steps_per_slide", 6))
    title_minimum_difference = float(policy.get("title_minimum_difference_pt", 5))
    title_minimum_ratio = float(policy.get("title_minimum_ratio", 1.2))
    presentation = Presentation(pptx_path)
    findings: list[dict[str, Any]] = []
    font_usage: Counter[str] = Counter()
    size_usage: Counter[float] = Counter()

    def add(slide: int, category: str, message: str, shape: str) -> None:
        findings.append(
            {
                "severity": "error",
                "slide": slide,
                "category": category,
                "message": message,
                "shape": shape,
            }
        )

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_bands: set[str] = set()
        title_sizes: list[float] = []
        non_title_sizes: list[float] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            shape_name = shape.name or ""
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    font = run.font.name
                    size = run.font.size.pt if run.font.size else None
                    if font:
                        font_usage[font] += len(run.text)
                        if allowed and font not in allowed:
                            add(
                                slide_number,
                                "font-family",
                                f"Font {font!r} is outside the allowed typography system.",
                                shape_name,
                            )
                    if size is None:
                        continue
                    value = round(float(size), 1)
                    size_usage[value] += len(run.text)
                    if value <= 13.5:
                        slide_bands.add("caption")
                    elif value <= 17.5:
                        slide_bands.add("body")
                    elif value <= 24:
                        slide_bands.add("emphasis")
                    elif value <= 34:
                        slide_bands.add("slide-title")
                    elif value >= 40:
                        slide_bands.add("deck-title")
                    else:
                        slide_bands.add("other")
                    if value < minimum:
                        add(
                            slide_number,
                            "minimum-font-size",
                            f"{value:g}pt is below the {minimum:g}pt minimum.",
                            shape_name,
                        )
                    if "deck-main-title" in shape_name and not in_range(value, deck_title_range):
                        add(
                            slide_number,
                            "deck-title-size",
                            f"Deck title {value:g}pt is outside {deck_title_range}.",
                            shape_name,
                        )
                    if shape_name == "slide-title" and not in_range(value, title_range):
                        add(
                            slide_number,
                            "slide-title-size",
                            f"Slide title {value:g}pt is outside {title_range}.",
                            shape_name,
                        )
                    if shape_name == "slide-title":
                        title_sizes.append(value)
                    elif "page-number" not in shape_name:
                        non_title_sizes.append(value)
                    if "page-number" in shape_name and not in_range(value, page_range):
                        add(
                            slide_number,
                            "page-number-size",
                            f"Page number {value:g}pt is outside {page_range}.",
                            shape_name,
                        )
        if len(slide_bands) > max_steps:
            add(
                slide_number,
                "size-step-count",
                (f"Slide uses {len(slide_bands)} typography bands; maximum is {max_steps}."),
                "slide",
            )
        if title_sizes and non_title_sizes:
            title_size = max(title_sizes)
            other_size = max(non_title_sizes)
            if (
                title_size - other_size < title_minimum_difference
                or title_size / other_size < title_minimum_ratio
            ):
                add(
                    slide_number,
                    "title-dominance",
                    (
                        f"Slide title is {title_size:g}pt but the largest other text "
                        f"is {other_size:g}pt; require at least "
                        f"{title_minimum_difference:g}pt and {title_minimum_ratio:g}x."
                    ),
                    "slide-title",
                )

    source_root = pptx_path.resolve().parent
    generated_sources = policy.get("generated_diagram_sources", [])
    generated_audit: list[dict[str, Any]] = []
    if isinstance(generated_sources, list):
        for raw_path in generated_sources:
            item = raw_path if isinstance(raw_path, dict) else {"path": raw_path}
            relative = Path(str(item.get("path", "")))
            source_path = source_root / relative
            if not source_path.exists():
                add(
                    0,
                    "generated-diagram-source",
                    f"Generated diagram source does not exist: {relative}",
                    str(relative),
                )
                continue
            text = source_path.read_text(encoding="utf-8")
            fonts = sorted(set(re.findall(r'fontname\s*=\s*"([^"]+)"', text)))
            sizes = sorted(
                {float(value) for value in re.findall(r"fontsize\s*=\s*([0-9.]+)", text)}
            )
            for font in fonts:
                if allowed and font not in allowed:
                    add(
                        0,
                        "generated-diagram-font",
                        f"Generated diagram {relative} uses unapproved font {font!r}.",
                        str(relative),
                    )
            apparent_sizes: list[float] = []
            image_value = item.get("image")
            placement_width = item.get("placement_width_inches")
            render_dpi = item.get("render_dpi")
            if image_value and placement_width and render_dpi and sizes:
                image_path = source_root / Path(str(image_value))
                if not image_path.exists():
                    add(
                        0,
                        "generated-diagram-image",
                        f"Generated diagram image does not exist: {image_value}",
                        str(relative),
                    )
                else:
                    image_width = Image.open(image_path).width
                    apparent_sizes = [
                        round(
                            size * float(render_dpi) * float(placement_width) / image_width,
                            2,
                        )
                        for size in sizes
                    ]
                    apparent_range = list(
                        policy.get("generated_diagram_apparent_font_range", [10, 18])
                    )
                    for size in apparent_sizes:
                        if not in_range(size, apparent_range):
                            add(
                                0,
                                "generated-diagram-apparent-font-size",
                                (
                                    f"Generated diagram {relative} appears at "
                                    f"{size:g}pt after placement; allowed range is "
                                    f"{apparent_range}."
                                ),
                                str(relative),
                            )
            generated_audit.append(
                {
                    "path": str(relative),
                    "fonts": fonts,
                    "source_font_sizes": sizes,
                    "apparent_font_sizes": apparent_sizes,
                }
            )

    errors = len(findings)
    return {
        "enabled": True,
        "passed": errors == 0,
        "errors": errors,
        "warnings": 0,
        "font_usage": dict(font_usage),
        "size_usage": {str(key): value for key, value in sorted(size_usage.items())},
        "generated_diagrams": generated_audit,
        "policy": policy,
        "findings": findings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit deck-wide font families and typography size hierarchy."
    )
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        report = audit(args.pptx, load_json(args.plan))
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        report = {
            "enabled": True,
            "passed": False,
            "errors": 1,
            "warnings": 0,
            "findings": [
                {
                    "severity": "error",
                    "slide": 0,
                    "category": "input",
                    "message": str(exc),
                }
            ],
        }
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
