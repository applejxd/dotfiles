from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from scripts.audit_slide_chrome import audit


def write_pptx(
    path: Path,
    *,
    header: str | None = None,
    footer: str | None = None,
    page_number: str | None = "01 / 01",
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(8), Inches(0.8)).text = "Title"
    if header:
        slide.shapes.add_textbox(Inches(0.7), Inches(0.1), Inches(4), Inches(0.3)).text = header
    if footer:
        slide.shapes.add_textbox(Inches(0.7), Inches(7.12), Inches(5), Inches(0.2)).text = footer
    if page_number:
        slide.shapes.add_textbox(
            Inches(11.8), Inches(7.12), Inches(0.8), Inches(0.2)
        ).text = page_number
    presentation.save(path)


def default_plan() -> dict:
    return {
        "meta": {
            "chrome_policy": {
                "enabled": True,
                "header": "none",
                "footer": "none",
                "page_numbers": "required",
            }
        }
    }


def test_page_number_only_passes(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    write_pptx(pptx)
    assert audit(pptx, default_plan())["passed"]


def test_default_header_fails(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    write_pptx(pptx, header="Section 1")
    report = audit(pptx, default_plan())
    assert not report["passed"]
    assert any(item["category"] == "unexpected-header" for item in report["findings"])


def test_default_footer_fails(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    write_pptx(pptx, footer="Source review")
    report = audit(pptx, default_plan())
    assert not report["passed"]
    assert any(item["category"] == "unexpected-footer" for item in report["findings"])


def test_missing_page_number_fails(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    write_pptx(pptx, page_number=None)
    report = audit(pptx, default_plan())
    assert not report["passed"]
    assert any(item["category"] == "missing-page-number" for item in report["findings"])


def test_user_requested_chrome_is_allowed(tmp_path: Path) -> None:
    pptx = tmp_path / "deck.pptx"
    write_pptx(pptx, header="Required brand header", footer="Required legal footer")
    plan = default_plan()
    plan["meta"]["chrome_policy"]["header"] = "user-requested"
    plan["meta"]["chrome_policy"]["footer"] = "user-requested"
    assert audit(pptx, plan)["passed"]
