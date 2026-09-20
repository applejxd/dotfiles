import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from scripts.validate_style import validate


def write_policy(path: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "surface_color": "F7F4EC",
                "base_color": "14213D",
                "main_color": "008F95",
                "accent_color": "E9A23B",
                "neutral_colors": ["FFFFFF", "5E6B78"],
                "allowed_background_colors": ["F7F4EC"],
                "max_background_colors": 1,
            }
        ),
        encoding="utf-8",
    )


def test_style_policy_accepts_declared_palette(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xF7, 0xF4, 0xEC)
    output = tmp_path / "deck.pptx"
    presentation.save(output)
    policy = tmp_path / "policy.json"
    write_policy(policy)
    assert validate(output, policy)["passed"]


def test_style_policy_rejects_extra_color(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    output = tmp_path / "deck.pptx"
    presentation.save(output)
    policy = tmp_path / "policy.json"
    write_policy(policy)
    assert not validate(output, policy)["passed"]
